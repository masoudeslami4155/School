#!/usr/bin/env python
"""4.59 regression: synthetic DB only, no live provider or private files."""
import io
import json
import unittest
from datetime import datetime
from unittest.mock import patch
from zoneinfo import ZoneInfo
import ai_agent_test as base
from flask import session, g
from openpyxl import load_workbook
from school_app import ai_workspace as memory
from school_app.ai_results import render_report, calculate, transform

A=base.AGENT
app=base.app

class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.ctx=app.test_request_context(); self.ctx.push()
        session.update(role='admin',user_id=501,personnel_number='ADMIN')
        g.ai_workspace=memory.new_workspace()
    def tearDown(self): self.ctx.pop()

    def test_normalize_full_name_arabic_digits(self):
        self.assertEqual(memory.normalize_fa('  عَلِي\u200cكريمي ۱۲۳ '),'علی کریمی 123')
        payload=json.loads(A.execute_tool('search_students',{'name':'علي نمونه'},'admin')['content'])
        self.assertEqual([r['code'] for r in payload],['AGT-1'])

    def test_paginated_results_have_total_and_no_overlap(self):
        a=A.execute_tool('search_students',{'page':1,'page_size':2},'admin')
        b=A.execute_tool('search_students',{'page':2,'page_size':2},'admin')
        self.assertEqual(a['meta']['total'],3); self.assertTrue(a['meta']['has_more'])
        self.assertFalse(b['meta']['has_more'])
        self.assertFalse({r['code'] for r in json.loads(a['content'])}&{r['code'] for r in json.loads(b['content'])})

    def test_jalali_context_nowruz_and_school_year(self):
        ctx=memory.temporal_context(datetime(2026,3,21,tzinfo=ZoneInfo('Asia/Tehran')))
        self.assertEqual(ctx['today'],'1405/01/01')
        self.assertEqual(ctx['academic_year'],'1404-1405')
        self.assertEqual(ctx['ranges']['ماه قبل'][0],'1404/12/01')
        self.assertEqual(ctx['ranges']['ماه قبل'][1],'1404/12/29')
        autumn=memory.temporal_context(datetime(2026,10,7,tzinfo=ZoneInfo('Asia/Tehran')))
        self.assertEqual(autumn['academic_year'],'1405-1406')

    def test_date_validation_and_attendance_range(self):
        self.assertEqual(memory.date_range('۱۴۰۵/۶/۱۲'),['1405/06/12']*2)
        with self.assertRaises(ValueError): memory.date_range('1405/13/01')
        response=A.execute_tool('get_attendance',{'date_from':'۱۴۰۵/۶/۱۲','date_to':'1405/06/12'},'admin')
        self.assertTrue(all(r['date']=='1405/06/12' for r in json.loads(response['content'])))

    def test_sql_limit_bounds_fetch_even_explicit_large_limit(self):
        payload=json.loads(A.execute_tool('query_database',{'sql':'WITH RECURSIVE n(x) AS (SELECT 1 UNION ALL SELECT x+1 FROM n WHERE x<1000) SELECT x FROM n LIMIT 1000'},'admin')['content'])
        self.assertEqual(payload['row_count'],200); self.assertTrue(payload['truncated'])

    def test_sql_timeout_and_readonly_failure(self):
        with patch.object(A,'_readonly_connection',return_value=None):
            with self.assertRaises(ValueError): A._fetch_rows('SELECT 1',[])
        with patch.object(A.time,'monotonic',side_effect=[0]+[100]*100):
            import sqlite3
            with self.assertRaises(sqlite3.OperationalError):
                A._fetch_rows('WITH RECURSIVE n(x) AS (SELECT 1 UNION ALL SELECT x+1 FROM n WHERE x<10000000) SELECT SUM(x) FROM n',[])

    def test_null_numeric_sql_parameters_preserved(self):
        data=json.loads(A.execute_tool('query_database',{'sql':'SELECT ? AS n, ? AS empty','params':[12,None]},'admin')['content'])
        self.assertEqual(data['rows'],[{'n':12,'empty':None}])

    def test_snapshot_report_escapes_and_keeps_zero(self):
        snap=memory.snapshot('test',{},[{'name':'<script>alert(1)</script>','amount':0}])
        text=render_report({'result_id':snap['result_id'],'title':'گزارش'})['content']
        self.assertIn('<td>0</td>',text); self.assertNotIn('<script>',text)
        self.assertIn('&lt;script&gt;',text)
        with self.assertRaises(ValueError): A.execute_tool('generate_report_html',{'title':'جعلی','headers':['a'],'rows':[['b']]},'admin')

    def test_decimal_arithmetic_does_not_claim_entire_population(self):
        snap=memory.snapshot('test',{},[{'amount':'0.1'},{'amount':'0.2'}],{'has_more':True})
        value=calculate({'result_id':snap['result_id'],'operation':'sum','column':'amount'})
        self.assertEqual(value['value'],'0.3'); self.assertIn('نه کل',value['coverage'])

    def test_transform_does_not_mutate_source(self):
        snap=memory.snapshot('test',{},[{'name':'B','x':10},{'name':'A','x':2}])
        rows,meta=transform({'result_id':snap['result_id'],'sort_by':'x'})
        self.assertEqual(rows[0]['x'],2); self.assertEqual(snap['rows'][0]['x'],10)
        self.assertEqual(meta['parent_result_id'],snap['result_id'])

    def test_save_reload_is_owned_and_permission_bound(self):
        state=memory.workspace(); memory.save(state)
        self.assertEqual(memory.load(state['id'])['id'],state['id'])
        session['user_id']=502
        with self.assertRaises(ValueError): memory.load(state['id'])
        session['user_id']=501; session['permissions']='__none__'
        with self.assertRaises(ValueError): memory.load(state['id'])

    def test_lease_and_atomic_delete(self):
        state=memory.workspace(); memory.save(state)
        lease=memory.acquire(state['id'])
        with self.assertRaises(ValueError): memory.acquire(state['id'])
        memory.release(lease)
        lease2=memory.acquire(state['id']); memory.release(lease2)
        memory.delete(state['id'])
        with self.assertRaises(ValueError): memory.load(state['id'])

    def test_memory_bounds(self):
        state=memory.workspace()
        for i in range(20): memory.snapshot('test',{},[{'i':i}])
        state['messages']=[{'role':'user','content':str(i)} for i in range(60)]
        memory.save(state); saved=memory.load(state['id'])
        self.assertEqual(len(saved['messages']),40); self.assertEqual(len(saved['results']),12)

    def test_blank_status_counts_match_every_tool(self):
        with base.get_db() as conn:
            conn.execute("INSERT INTO students(code,first_name,last_name,status) VALUES('BLANK-459','نام','نمونه','   ')")
        try:
            stats=json.loads(A.execute_tool('get_statistics',{'level':'grade'},'admin')['content'])
            summary=json.loads(A.execute_tool('get_school_summary',{'scope':'active'},'admin')['content'])
            search=A.execute_tool('search_students',{'status':'فعال'},'admin')
            self.assertEqual(sum(r['count'] for r in stats),summary['total_students'])
            self.assertEqual(search['meta']['total'],summary['total_students'])
        finally:
            with base.get_db() as conn: conn.execute("DELETE FROM students WHERE code='BLANK-459'")

    def test_stream_is_lazy_and_preserves_actual_rows(self):
        seen=[]
        def provider(config,messages,tools=None):
            seen.append(json.loads(json.dumps(messages)))
            if len(seen)==1: return {'content':None,'tool_calls':[{'id':'x','type':'function','function':{'name':'search_students','arguments':'{}'}}]}
            return {'content':'تمام','tool_calls':[]}
        with patch.object(A,'_load_provider_config',return_value={'configured':True,'provider':'ollama'}),patch.object(A,'request_chat_message',side_effect=provider):
            events=A.iter_agent_events('فهرست',[])
            self.assertEqual(seen,[])
            first=next(events); self.assertEqual(first['event'],'tool_start'); self.assertEqual(len(seen),1)
            list(events)
        self.assertIn('AGT-1',seen[1][-1]['content'])
        self.assertIn('result_id',seen[1][-1]['content'])

    def test_multiple_tool_calls_keep_protocol_pairing(self):
        calls=[{'id':f'c{i}','type':'function','function':{'name':'get_statistics','arguments':json.dumps({'level':level})}} for i,level in enumerate(['grade','gender'])]
        with patch.object(A,'_load_provider_config',return_value={'configured':True,'provider':'ollama'}),patch.object(A,'request_chat_message',side_effect=[{'content':None,'tool_calls':calls},{'content':'تمام','tool_calls':[]}]):
            messages,_,_=A.run_agent('آمار')
        assistant=[m for m in messages if m.get('tool_calls')]
        self.assertEqual(len(assistant),1); self.assertEqual(len(assistant[0]['tool_calls']),2)
        self.assertEqual(len([m for m in messages if m['role']=='tool']),2)

class RouteTests(unittest.TestCase):
    def setUp(self):
        self.client=base.page_client()
        base.set_session(self.client,'admin',user_id=701)
        base.save_provider()
    def post(self,data):
        return self.client.post('/ai-agent/chat',json=data,headers={'X-CSRF-Token':base.CSRF})
    def test_post_history_and_persistent_reload(self):
        seen=[]
        def provider(config,messages,tools=None):
            seen.append(json.loads(json.dumps(messages))); return {'content':'پاسخ آزمایشی','tool_calls':[]}
        with patch.object(A,'request_chat_message',side_effect=provider):
            first=self.post({'query':'اول','history':[{'role':'user','content':'زمینه قبلی'}]}).get_json()
            second=self.post({'query':'دوم','conversation_id':first['conversation_id']})
        self.assertEqual(second.status_code,200)
        self.assertIn('زمینه قبلی',str(seen[0]))
        self.assertIn('اول',str(seen[1]))
        rv=self.client.get('/ai-agent/chat?conversation_id='+first['conversation_id'])
        self.assertEqual(len(rv.get_json()['conversation']['messages']),4)
        self.assertEqual(self.post({'action':'rename','conversation_id':first['conversation_id'],'title':'جدید'}).status_code,200)
        self.assertEqual(self.post({'action':'delete','conversation_id':first['conversation_id']}).status_code,200)
        self.assertEqual(self.client.get('/ai-agent/chat?conversation_id='+first['conversation_id']).status_code,400)

    def test_post_stream_csrf_and_persisted_sources(self):
        self.assertEqual(self.client.post('/ai-agent/stream',json={'query':'سلام'}).status_code,400)
        call={'id':'x','type':'function','function':{'name':'search_students','arguments':'{"page_size":1}'}}
        with patch.object(A,'request_chat_message',side_effect=[{'content':None,'tool_calls':[call]},{'content':'یک صفحه','tool_calls':[]}]):
            response=self.client.post('/ai-agent/stream',json={'query':'فهرست'},headers={'X-CSRF-Token':base.CSRF})
            events=[json.loads(x[6:]) for x in response.get_data(as_text=True).split('\n\n') if x.startswith('data: ')]
        self.assertEqual(events[-1]['event'],'done')
        cid=next(x['data']['id'] for x in events if x['event']=='conversation')
        rid=next(x['data']['result_id'] for x in events if x['event']=='source')
        loaded=self.client.get('/ai-agent/chat?conversation_id='+cid).get_json()['conversation']
        self.assertIn(rid,loaded['results']); self.assertTrue(loaded['messages'][1]['sources'])
        exported=self.post({'action':'export','conversation_id':cid,'result_id':rid})
        self.assertEqual(exported.status_code,200)
        wb=load_workbook(io.BytesIO(exported.data)); self.assertEqual(wb.active.max_row,2)
        allrows=self.post({'action':'export','conversation_id':cid,'result_id':rid,'all':True})
        self.assertEqual(allrows.status_code,200)
        wb=load_workbook(io.BytesIO(allrows.data)); self.assertEqual(wb.active.max_row,4)
        other=base.page_client(); base.set_session(other,'admin',user_id=702)
        self.assertEqual(other.get('/ai-agent/chat?conversation_id='+cid).status_code,400)

    def test_stream_disconnect_releases_lease(self):
        cid=self.post({'action':'new'}).get_json()['conversation_id']
        response=self.client.post('/ai-agent/stream',json={'query':'سلام','conversation_id':cid},
                                  headers={'X-CSRF-Token':base.CSRF},buffered=False)
        self.assertEqual(response.status_code,200)
        response.close()  # browser closes after receiving start, before provider call
        self.assertEqual(self.post({'action':'rename','conversation_id':cid,'title':'آزاد'}).status_code,200)

    def test_stream_provider_failure_is_structured_and_retryable(self):
        cid=self.post({'action':'new'}).get_json()['conversation_id']
        with patch.object(A,'request_chat_message',side_effect=A.AIProviderError('قطع آزمایشی')):
            response=self.client.post('/ai-agent/stream',json={'query':'سلام','conversation_id':cid},headers={'X-CSRF-Token':base.CSRF})
            payload=response.get_data(as_text=True)
        self.assertIn('قطع آزمایشی',payload)
        self.assertNotIn('"event": "done"',payload)
        self.assertEqual(self.post({'action':'rename','conversation_id':cid,'title':'آزاد'}).status_code,200)

    def test_malformed_requests_and_provider_failure_release_lease(self):
        self.assertEqual(self.post([]).status_code,400)
        cid=self.post({'action':'new'}).get_json()['conversation_id']
        with patch.object(A,'request_chat_message',side_effect=A.AIProviderError('آزمایشی')):
            rv=self.post({'conversation_id':cid,'query':'سلام'})
        self.assertEqual(rv.status_code,502)
        self.assertEqual(self.post({'action':'rename','conversation_id':cid,'title':'آزاد'}).status_code,200)

if __name__=='__main__': unittest.main(verbosity=2)
