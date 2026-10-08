#!/usr/bin/env python
import io
import unittest
from urllib.parse import urlencode
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from flask import template_rendered
import statistics_fixture as F
from school_app.statistics_routes import _grade_rank

class FrequencyTests(unittest.TestCase):
    def setUp(self):self.client=F.client()
    def get(self,query=None,client=None):
        contexts=[]
        def collect(sender,template,context,**kw):contexts.append(context)
        url=query if isinstance(query,str) else '/statistics'+('?' + urlencode(query,doseq=True) if query else '')
        with template_rendered.connected_to(collect,F.app):r=(client or self.client).get(url)
        return r,contexts[-1] if contexts else {}
    def test_totals_no_double_count_from_teacher_whitespace_duplicates(self):
        r,c=self.get();self.assertEqual(r.status_code,200)
        self.assertEqual(c['population'],10);self.assertEqual(c['frequency_totals']['total'],10)
        self.assertEqual([c['frequency_totals'][k] for k in ('boys','girls','missing')],[5,3,2])
        self.assertEqual(c['teacher_count'],2)
    def test_exact_preparation_grade_order(self):
        self.assertLess(_grade_rank('آمادگی'),_grade_rank('آمادگی مقدماتی'))
        self.assertLess(_grade_rank('آمادگی مقدماتی'),_grade_rank('آمادگی تکمیلی'))
        self.assertLess(_grade_rank('پایه اول'),_grade_rank('پایه دوم'))
    def test_grade_subtotal_labels_not_repeated(self):
        r,c=self.get();labels=[x['label'] for x in c['frequency_display_rows'] if x['row_kind']=='grade_total']
        self.assertIn('جمع پایه اول',labels);self.assertNotIn('جمع پایه پایه اول',labels)
    def test_three_views_same_population(self):
        for view,key in [('detail',None),('grade','grade'),('teacher','teacher')]:
            r,c=self.get({'view':view});self.assertEqual(r.status_code,200)
            self.assertEqual(sum(x['total'] for x in c['frequency_rows']),10)
            if key:
                self.assertEqual(len(c['frequency_rows']),5 if key=='grade' else 3)
                self.assertTrue(all(x['row_kind']=='detail' for x in c['frequency_display_rows']))
                self.assertNotIn('teacher' if key=='grade' else 'grade',[col['key'] for col in c['frequency_columns_view']])
    def test_subtotal_toggles_are_independent(self):
        for grade,teacher in [(0,0),(1,0),(0,1),(1,1)]:
            _,c=self.get({'grade_totals':grade,'teacher_totals':teacher})
            kinds=[x['row_kind'] for x in c['frequency_display_rows']]
            self.assertEqual('grade_total' in kinds,bool(grade));self.assertEqual('teacher_total' in kinds,bool(teacher))
            self.assertEqual(c['frequency_totals']['total'],10)
    def test_empty_explicit_selection_is_not_all(self):
        for missing in ('print_teacher','print_grade','print_col'):
            q={'selection':1,'print':1,'print_teacher':['T1'],'print_grade':['پایه اول'],'print_col':['grade','total']};q.pop(missing)
            r,_=self.get(q);self.assertEqual(r.status_code,400)
    def test_legacy_print_default_still_works(self):
        r,c=self.get({'print':1});self.assertEqual(r.status_code,200);self.assertEqual(c['selected_population'],10)
    def test_numeric_only_explicit_columns_rejected(self):
        r,_=self.get({'selection':1,'print_teacher':'T1','print_grade':'پایه اول','print_col':['total','share']});self.assertEqual(r.status_code,400)
    def test_filtered_print_recalculates_totals_but_keeps_denominator(self):
        _,c=self.get({'selection':1,'print':1,'print_teacher':['T1'],'print_grade':['پایه اول'],'print_col':['grade','total','share']})
        self.assertEqual(c['frequency_totals']['total'],3);self.assertEqual(c['population'],10);self.assertEqual(c['frequency_totals']['share'],30.0)
        self.assertEqual([x['total'] for x in c['frequency_display_rows'] if x['row_kind']=='teacher_total'],[3])
    def test_every_cell_drills_into_its_exact_count(self):
        for view in ('detail','grade','teacher'):
            _,c=self.get({'view':view,'print_grade':'پایه اول'})
            for row in c['frequency_display_rows']+[c['frequency_totals']]:
                for key,url in row['links'].items():
                    r,detail=self.get(url);self.assertEqual(r.status_code,200)
                    self.assertEqual(detail['detail_count'],row[key],(view,key,url))
    def test_active_and_additional_filters(self):
        _,c=self.get({'status':'active'});self.assertEqual(c['population'],8)
        _,c=self.get({'class_name':'الف','sida_class':'سیدا الف'});self.assertEqual(c['population'],4)
        self.assertEqual(len(c['active_chips']),2)
        _,cleared=self.get(c['active_chips'][0]['url']);self.assertEqual(len(cleared['active_chips']),1)
    def test_quality_warning_counts_and_links(self):
        _,c=self.get();expected={'no_teacher':1,'invalid_teacher':1,'missing_gender':2,'missing_grade':1}
        self.assertEqual({x['key']:x['count'] for x in c['quality_items']},expected)
        for item in c['quality_items']:
            _,detail=self.get(item['url']);self.assertEqual(detail['detail_count'],item['count'])
        _,c=self.get({'class_name':'الف'});self.assertEqual(c['quality_items'],[])
    def test_role_and_permission_boundaries(self):
        for query in ('','?print=1','?action=students','?action=export&format=xlsx'):
            self.assertEqual(F.client('teacher').get('/statistics'+query).status_code,403)
            self.assertEqual(F.client('manager','student_profile').get('/statistics'+query).status_code,403)
        restricted=F.client('manager','statistics')
        r,c=self.get(client=restricted);self.assertEqual(r.status_code,200);self.assertFalse(c['can_drill'])
        self.assertTrue(all(not row['links'] for row in c['frequency_display_rows']))
        self.assertEqual(restricted.get('/statistics?action=students').status_code,403)
        partial=F.client('manager','statistics,student_profile')
        self.assertEqual(partial.get('/statistics?action=students').status_code,200)
        self.assertEqual(partial.get('/statistics?action=students&format=xlsx').status_code,403)
    def test_summary_excel_matches_selected_rows_and_numeric_types(self):
        q={'action':'export','format':'xlsx','selection':1,'view':'grade','print_teacher':['T1','T2'],'print_grade':['پایه اول'],'print_col':['grade','boys','girls','total','share']}
        r,_=self.get(q);self.assertEqual(r.status_code,200)
        ws=load_workbook(io.BytesIO(r.data)).active
        self.assertEqual([c.value for c in ws[5]],['پایه','پسر','دختر','جمع','سهم'])
        self.assertEqual([c.value for c in ws[6]],['پایه اول',2,2,4,.4]);self.assertEqual(ws['D6'].data_type,'n')
        self.assertTrue(ws.sheet_view.rightToLeft);self.assertEqual(ws['E6'].number_format,'0.0%')
    def test_student_excel_keeps_codes_and_is_full_filtered_list(self):
        r,_=self.get({'action':'students','drill_teacher':'T1','drill_grade':'پایه اول','metric':'boys','format':'xlsx'})
        ws=load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws['B6'].value,'0001');self.assertEqual(ws['B7'].value,'0002');self.assertEqual(ws['B6'].data_type,'s')
    def test_formula_like_name_is_not_executable_in_excel(self):
        with F.app.app_context():
            with F.get_db() as c:c.execute("UPDATE students SET first_name=? WHERE code='0001'",('=SUM(1,2)',))
        try:
            r,_=self.get({'action':'students','drill_grade':'پایه اول','metric':'boys','format':'xlsx'})
            ws=load_workbook(io.BytesIO(r.data)).active
            cell=next(c for row in ws for c in row if isinstance(c.value,str) and c.value.startswith('=SUM'))
            self.assertEqual(cell.data_type,'s')
        finally:
            with F.app.app_context():
                with F.get_db() as c:c.execute("UPDATE students SET first_name='نمونه 1' WHERE code='0001'")
    def test_empty_scope_and_missing_grade_drill(self):
        r,c=self.get({'grade':'غیرموجود'});self.assertEqual(r.status_code,200);self.assertEqual(c['selected_population'],0)
        _,c=self.get({'grade':'ثبت نشده'});self.assertEqual(c['population'],1)
        _,d=self.get({'action':'students','drill_grade':'ثبت نشده'});self.assertEqual(d['detail_count'],1)
    def test_invalid_query_returns_clear_error(self):
        for q in ({'view':'unsafe'},{'action':'bad'},{'action':'students','quality':'sql'},{'action':'students','page':'x'},{'orientation':'anything'},{'action':'export','format':'exe'}):
            self.assertEqual(self.get(q)[0].status_code,400)
    def test_detail_pagination_and_full_export(self):
        with F.app.app_context():
            with F.get_db() as c:
                for i in range(55):c.execute("INSERT INTO students(code,first_name,last_name,status,grade,class_name,teacher_code) VALUES(?,?,?,'فعال','پایه اول','PAGE-CHECK','T1')",(f'PAGE-{i:03d}','نمونه',str(i)))
        try:
            _,first=self.get({'action':'students','class_name':'PAGE-CHECK'})
            self.assertEqual(first['detail_count'],55);self.assertEqual(len(first['student_rows']),50)
            _,second=self.get(first['next_url']);self.assertEqual(len(second['student_rows']),5)
            self.assertFalse({x['id'] for x in first['student_rows']} & {x['id'] for x in second['student_rows']})
            r,_=self.get(second['detail_export_urls']['xlsx'])
            sheet=load_workbook(io.BytesIO(r.data)).active
            self.assertEqual(sheet.max_row,61)
        finally:
            with F.app.app_context():
                with F.get_db() as c:c.execute("DELETE FROM students WHERE class_name='PAGE-CHECK'")

    def test_request_cannot_escape_source_filter(self):
        _,c=self.get({'action':'students','teacher_code':'T1','drill_teacher':'T2'})
        self.assertEqual(c['detail_count'],0)
        r,_=self.get({'teacher_code':"T1' OR 1=1 --"})
        self.assertEqual(r.status_code,200)
        self.assertIn(b'private, no-store',r.headers['Cache-Control'].encode())

    def test_quality_counts_respect_selected_output(self):
        _,c=self.get({'print_teacher':['T1'],'print_grade':['پایه دوم']})
        self.assertEqual({x['key']:x['count'] for x in c['quality_items']},{'missing_gender':1})
        _,d=self.get(c['quality_items'][0]['url']);self.assertEqual(d['detail_count'],1)

    def test_reading_stats_does_not_modify_source(self):
        with F.app.app_context():
            with F.get_db() as c:before=[tuple(r) for r in c.execute('SELECT * FROM students ORDER BY id')]
        self.get();self.get({'action':'students'});self.get({'action':'export','format':'xlsx'})
        with F.app.app_context():
            with F.get_db() as c:after=[tuple(r) for r in c.execute('SELECT * FROM students ORDER BY id')]
        self.assertEqual(before,after)

if __name__=='__main__':unittest.main(verbosity=2)
