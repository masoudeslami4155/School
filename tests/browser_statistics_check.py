#!/usr/bin/env python
"""Real Chromium: all seven statistics improvements with synthetic data."""
import io
from threading import Thread
from pathlib import Path
from werkzeug.serving import make_server
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
import statistics_fixture as F
client=F.client();cookie=client.get_cookie('session').value
server=make_server('127.0.0.1',0,F.app,threaded=True);Thread(target=server.serve_forever,daemon=True).start()
url=f'http://127.0.0.1:{server.server_port}'
try:
 with sync_playwright() as p:
  browser=p.chromium.launch(headless=True)
  page=browser.new_page(viewport={'width':1440,'height':1100});errors=[]
  page.context.add_cookies([{'name':'session','value':cookie,'url':url}]);page.on('pageerror',lambda e:errors.append(str(e)))
  page.goto(url+'/statistics');page.locator('#frequencyTable').wait_for()
  assert page.locator('#selectedPopulation').inner_text()=='10'
  assert page.locator('.stats-quality a').count()==4
  assert not page.locator('#statisticsPrintOptions').get_attribute('open')
  assert 'جمع پایه پایه' not in page.locator('#frequencyTable').inner_text()
  page.screenshot(path='/home/user/statistics-463-desktop.png',full_page=True)
  # Every interactive number opens a precise list and preserves return scope.
  href=page.locator('#frequencyTable tr.detail-row a.stats-count').first.get_attribute('href')
  count=page.locator('#frequencyTable tr.detail-row a.stats-count').first.inner_text()
  page.goto(url+href);assert count+' دانش‌آموز' in page.locator('.stats-hero h1').inner_text()
  assert page.locator('#statisticsStudentTable tbody tr').count()==int(count)
  with page.expect_download() as d:page.locator('[data-student-export="xlsx"]').click()
  ws=load_workbook(io.BytesIO(Path(d.value.path()).read_bytes())).active;assert ws.max_row==int(count)+6
  page.get_by_role('link',name='بازگشت به همان آمار').click()
  # Independent subtotals and three render modes.
  page.locator('input[type=checkbox][name=grade_totals]').uncheck();page.locator('input[type=checkbox][name=teacher_totals]').uncheck()
  page.get_by_role('button',name='اعمال نمایش',exact=True).click();assert page.locator('#frequencyTable .subtotal-row').count()==0
  for view,rows in [('grade',5),('teacher',3),('detail',7)]:
   page.locator('#statisticsView').select_option(view);page.get_by_role('button',name='اعمال نمایش',exact=True).click()
   page.locator('#frequencyTable').wait_for()
   assert page.locator('#frequencyTable tr.detail-row').count()==rows,(view,page.locator('#frequencyTable tr.detail-row').count(),page.url)
  # More filters + removable chips.
  page.locator('.stats-more summary').click();page.locator('#class_name').select_option('الف')
  page.get_by_role('button',name='اعمال فیلتر',exact=True).click();assert page.locator('#selectedPopulation').inner_text()=='4'
  assert page.locator('.stats-quality').count()==0
  page.locator('.stats-chip').first.click();assert page.locator('#selectedPopulation').inner_text()=='10'
  page.locator('input[type=checkbox][name=grade_totals]').check();page.locator('input[type=checkbox][name=teacher_totals]').check();page.get_by_role('button',name='اعمال نمایش',exact=True).click()
  page.locator('#statisticsPrintOptions > summary').click()
  page.locator('[data-select-group="teacher"][data-select-state="none"]').click()
  page.get_by_role('button',name='چاپ انتخاب‌ها',exact=True).click()
  assert 'حداقل یک معلم' in page.locator('#statisticsExportStatus').inner_text()
  page.locator('[data-group="teacher"][value="T1"]').check()
  page.locator('[data-select-group="grade"][data-select-state="none"]').click();page.locator('[data-group="grade"][value="پایه اول"]').check()
  assert '3 دانش‌آموز' in page.locator('#statisticsSelectionPreview').inner_text()
  page.locator('#statisticsOrientation').select_option('landscape')
  for kind in ['xlsx','pdf']:
   with page.expect_download(timeout=45000) as d:page.locator('[data-stats-export="'+kind+'"]').click()
   assert d.value.suggested_filename=='statistics.'+kind
   if kind=='pdf':d.value.save_as('/home/user/statistics-463-sample.pdf')
   else:
    ws=load_workbook(io.BytesIO(Path(d.value.path()).read_bytes())).active
    assert ws.page_setup.orientation=='landscape'
    assert all('دوم' not in str(cell.value) for row in list(ws.rows)[5:] for cell in row)
  with page.expect_popup() as popup:page.get_by_role('button',name='چاپ انتخاب‌ها',exact=True).click()
  printed=popup.value;printed.wait_for_load_state();assert printed.locator('#selectedPopulation').inner_text()=='3'
  assert printed.locator('#frequencyTable th').count()>0
  printed.evaluate('document.fonts.ready');printed.pdf(path='/home/user/statistics-463-browser-print.pdf',prefer_css_page_size=True,print_background=True);printed.close()
  page.set_viewport_size({'width':390,'height':844});page.goto(url+'/statistics');page.locator('#frequencyTable').wait_for()
  assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+2'), page.evaluate("[...document.querySelectorAll('body *')].map(e=>({id:e.id,tag:e.tagName,cls:String(e.className),width:e.getBoundingClientRect().width,right:e.getBoundingClientRect().right,left:e.getBoundingClientRect().left})).filter(x=>x.width>innerWidth||x.left<0).slice(0,22)")
  page.screenshot(path='/home/user/statistics-463-mobile.png',full_page=True)
  page.evaluate("document.documentElement.dataset.theme='dark'")
  assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+2')
  page.screenshot(path='/home/user/statistics-463-mobile-dark.png',full_page=True)
  assert not errors,errors
  browser.close();print('PASS: filters, 3 views, subtotals, exact drilldown, quality warnings, empty selections, Excel/PDF, print popup, mobile/light/dark, no JS errors')
finally:server.shutdown()
