#!/usr/bin/env python
"""Isolated 4.62 design/export regression; provider is synthetic."""
import io
import json
import unittest
from unittest.mock import patch
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
import ai_agent_test as base
from school_app import ai_design as design
from school_app import ai_document_store as archive
from school_app.database import get_db
from flask import session

class DesignTests(unittest.TestCase):
    def setUp(self):
        self.ctx=base.app.app_context();self.ctx.push();base.save_provider()
        self.client=base.page_client();base.set_session(self.client,'admin',user_id=862)
        self.other=base.page_client();base.set_session(self.other,'manager',user_id=863)
        with get_db() as conn:
            conn.execute("DELETE FROM app_settings WHERE key LIKE 'ai_document_archive:%' OR key LIKE 'ai_control:%'")
    def tearDown(self):self.ctx.pop()
    def post(self,data,client=None):
        return (client or self.client).post('/ai-documents',json=data,headers={'X-CSRF-Token':base.CSRF})
    def generate(self,kind='table'):
        data={'action':'design_generate','kind':kind,'title':'گزارش مدرن','body':'کد\tنام\n0012\tآزمایشی\n0013\t=SUM(A1:A2)' if kind=='table' else 'با سلام\nبرای بررسی پرونده مراجعه فرمایید.','design_brief':'طراحی آبی رسمی'}
        with patch('school_app.ai_documents.request_chat_completion',return_value=json.dumps({'html':design.DEFAULT_LAYOUT})) as model:
            response=self.post(data)
            self.assertEqual(response.status_code,200,response.data)
            prompt=str(model.call_args)
            self.assertNotIn('0012',prompt)
            self.assertNotIn('آزمایشی',prompt)
        return response.json['document']
    def test_generate_save_and_owner_isolation(self):
        doc=self.generate();self.assertEqual(doc['document_type'],'designed_table')
        response=self.client.get('/ai-documents?api=1&id='+doc['id'])
        self.assertEqual(response.json['document']['html'],doc['html'])
        self.assertEqual(self.other.get('/ai-documents?api=1&id='+doc['id']).status_code,400)
    def test_payload_contains_design_only(self):
        r=self.post({'action':'design_preview','kind':'letter','title':'محرمانه','body':'نام نباید ارسال شود','design_brief':'آبی'})
        self.assertEqual(r.status_code,200)
        self.assertNotIn('نام نباید',json.dumps(r.json,ensure_ascii=False))
    def test_real_word_tables_and_safe_excel(self):
        doc=self.generate()
        word=self.post({'action':'export','id':doc['id'],'format':'docx'})
        d=Document(io.BytesIO(word.data));self.assertEqual(len(d.tables),1)
        self.assertEqual(d.tables[0].cell(1,0).text,'0012')
        excel=self.post({'action':'export','id':doc['id'],'format':'xlsx'})
        sheet=load_workbook(io.BytesIO(excel.data)).active
        self.assertEqual(sheet['A3'].value,'0012')
        self.assertEqual(sheet['B4'].value,'=SUM(A1:A2)')
        self.assertEqual(sheet['B4'].data_type,'s')
        self.assertTrue(sheet.sheet_view.rightToLeft)
    def test_letter_excel_is_explicit_error(self):
        doc=self.generate('letter')
        r=self.post({'action':'export','id':doc['id'],'format':'xlsx'})
        self.assertEqual(r.status_code,400);self.assertIn('جدول ندارد',r.json['error'])
    def test_styles_and_active_content_safety(self):
        clean=design.clean('<script>alert(1)</script><style>body{}</style><p onclick="x()" style="color:#123456;background:url(https://example.invalid)">متن</p><img src="x">')
        self.assertNotIn('script',clean);self.assertNotIn('alert',clean);self.assertNotIn('onclick',clean)
        self.assertNotIn('https:',clean);self.assertNotIn('<img',clean)
        with self.assertRaises(ValueError):design.compose('<div style="color:{{content}}">bad</div>','<p>متن</p>')
        with self.assertRaises(ValueError):design.compose('<div>{{content}}{{content}}</div>','<p>متن</p>')
        with self.assertRaises(ValueError):design.compose('<h1>آمار ساختگی 999</h1><div>{{content}}</div>','<p>واقعی</p>')
    def test_merged_cells(self):
        doc={'html':'<table><tr><th colspan="2">عنوان</th></tr><tr><td>001</td><td>دو</td></tr></table>','title':'آزمون','status':'draft','paper':'a4','number':''}
        sheet=load_workbook(design.excel(doc)).active
        self.assertIn('A2:B2',[str(r) for r in sheet.merged_cells.ranges])
        d=Document(design.word(doc));self.assertEqual(d.tables[0].cell(0,1).text,'عنوان')
    def test_invalid_model_does_not_save(self):
        with patch('school_app.ai_documents.request_chat_completion',return_value='{"html":"<p>بدون جای‌نگهدار</p>"}'):
            response=self.post({'action':'design_generate','body':'متن','kind':'letter'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.client.get('/ai-documents?api=1').json['documents'],[])
    def test_agent_letter_tool_archives_and_streams(self):
        def provider(config,messages,tools=None):
            if not any(m.get('role')=='tool' for m in messages):
                return {'content':None,'tool_calls':[{'id':'design','type':'function','function':{'name':'design_letter','arguments':json.dumps({'title':'دعوت‌نامه','html':'<h1>دعوت‌نامه</h1><p style="color:#24576b">با سلام، برای جلسه مراجعه کنید.</p><p>مهر و امضا</p>'})}}]}
            return {'content':'پیش‌نویس ذخیره شد.','tool_calls':[]}
        with patch.object(base.AGENT,'request_chat_message',side_effect=provider):
            r=self.client.post('/ai-agent/stream',json={'query':'نامه‌ای مدرن برای دعوت به جلسه بساز'},headers={'X-CSRF-Token':base.CSRF},buffered=True)
        self.assertEqual(r.status_code,200)
        self.assertIn('data-archive-id',r.get_data(as_text=True))
        self.assertIn('"event": "report"',r.get_data(as_text=True))
        docs=self.client.get('/ai-documents?api=1').json['documents'];self.assertEqual(len(docs),1)
    def test_graphic_html_is_self_contained(self):
        doc=self.generate()
        r=self.post({'action':'export','id':doc['id'],'format':'html'})
        body=r.data.decode();self.assertIn('data:font/ttf;base64,',body)
        self.assertIn('0012',body);self.assertNotIn('<script',body)
    def test_remote_generation_requires_consent(self):
        with get_db() as conn:
            base.save_provider_config(conn,{'provider':'openai-compatible','base_url':'https://example.invalid/v1','model':'fixture','api_key':'synthetic-test-key'})
        with patch('school_app.ai_documents.request_chat_completion',return_value=json.dumps({'html':design.DEFAULT_LAYOUT})) as model:
            r=self.post({'action':'design_generate','kind':'letter','body':'خصوصی'})
            self.assertEqual(r.status_code,400)
            model.assert_not_called()
            r=self.post({'action':'design_generate','kind':'letter','body':'خصوصی','allow_external_data':True})
            self.assertEqual(r.status_code,200,r.data)
            self.assertNotIn('خصوصی',str(model.call_args))

    def test_agent_design_keeps_real_rows_and_stores_source(self):
        with base.app.test_request_context():
            session.update({'role':'admin','user_id':862})
            result=base.report_fixture()
            rendered=base.AGENT.execute_tool('generate_report_html',{'result_id':result,'title':'گزارش','layout_html':'<div style="border:2px solid #123456">{{content}}</div>'},'admin')['content']
            soup=BeautifulSoup(rendered,'html.parser')
            self.assertEqual([c.get_text() for c in soup.select('tbody td')],['علی نمونه','حاضر','مریم نمونه','غایب'])
            ident=soup.select_one('[data-archive-id]')['data-archive-id']
            with get_db() as conn:doc=archive.load(conn,ident)
            self.assertIn(result,doc['html'])
            self.assertIn('border:2px solid #123456',doc['html'])

    def test_edit_revision_and_restore(self):
        doc=self.generate('letter')
        r=self.post({'action':'edit','id':doc['id'],'revision':1,'html':doc['html']+'<p>اصلاح</p>'})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['document']['revision'],2)
        restored=self.post({'action':'restore','id':doc['id'],'revision':2,'target':1})
        self.assertEqual(restored.json['document']['html'],doc['html'])

if __name__=='__main__':unittest.main(verbosity=2)
