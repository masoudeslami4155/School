from __future__ import annotations

from datetime import datetime
from io import BytesIO

from flask import current_app, flash, redirect, render_template, request, send_file, session, url_for

from ..database import auto_backup, get_db
from ..dates import (MONTH_NAMES, get_today_jalali, is_jalali_leap, normalize_jalali_date,
                     parse_jalali_date, today_string)
from ..print_layouts import load_layout, prepare_for_print
from ..report_print_layouts import load_layout as load_report_layout
from ..security import audit

_STATUS_COUNTER_KEYS = {'حاضر': 'present', 'غایب': 'absent', 'مرخصی': 'excused', 'تاخیر': 'late'}

_STATUSES = ('حاضر', 'غایب', 'مرخصی', 'تاخیر')
_PER_PAGE_OPTIONS = (25, 50, 100, 200)
_SORT_COLUMNS = {
    'date': 'a.date',
    'name': 's.last_name',
    'class': 's.class_name',
    'grade': 's.grade',
    'status': 'a.status',
}
_RANGE_LABELS = {
    'today': 'امروز',
    'week': '۷ روز اخیر',
    'month': 'این ماه',
    'last_month': 'ماه گذشته',
    'term': 'سال تحصیلی',
    'all': 'همه بازه',
    'custom': 'بازه دلخواه',
}
_WEEKDAY_NAMES = ('شنبه', 'یک‌شنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنج‌شنبه', 'جمعه')
_JALALI_LEAPS = (1, 5, 9, 13, 17, 21, 25, 29)
_ANALYTICS_COLUMNS = ('code', 'first_name', 'last_name', 'class_name', 'grade', 'date', 'status')

SIDA_CLASS_OPTIONS = [
    'اول 1 استثنایی', 'اول 2 استثنایی', 'اول 3 استثنایی',
    'دوم ابتدایی استثنایی', 'سوم ابتدایی استثنایی', 'چهارم ابتدایی استثنایی',
    'پنجم ابتدایی استثنایی', 'ششم ابتدایی استثنایی',
    'آمادگی مقدماتی', 'آمادگی تکمیلی',
]


def _group_records_by_month(records):
    """Group flat attendance rows into per-month sections for a clear, presentable report."""
    groups: "dict[str, dict]" = {}
    order: list[str] = []
    for row in records:
        date_val = (row['date'] or '').strip()
        month_key = date_val[:7] if len(date_val) >= 7 and '/' in date_val else 'بدون تاریخ معتبر'
        if month_key not in groups:
            label = month_key
            parts = date_val.split('/')
            if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit() and 1 <= int(parts[1]) <= 12:
                label = f"{MONTH_NAMES[int(parts[1]) - 1]} {parts[0]}"
            groups[month_key] = {
                'key': month_key, 'label': label, 'rows': [],
                'present': 0, 'absent': 0, 'excused': 0, 'late': 0,
            }
            order.append(month_key)
        grp = groups[month_key]
        grp['rows'].append(row)
        counter_key = _STATUS_COUNTER_KEYS.get(row['status'])
        if counter_key:
            grp[counter_key] += 1
    return [groups[key] for key in order]


def _jalali_day_number(year: int, month: int, day: int) -> int:
    def leaps_upto(value: int) -> int:
        if value <= 0:
            return 0
        cycles, remainder = divmod(value, 33)
        return cycles * len(_JALALI_LEAPS) + sum(1 for leap in _JALALI_LEAPS if leap <= remainder)

    total = (year - 1) * 365 + leaps_upto(year - 1)
    for current in range(1, month):
        if current <= 6:
            total += 31
        elif current <= 11:
            total += 30
        else:
            total += 30 if is_jalali_leap(year) else 29
    return total + day


def _jalali_from_day_number(number: int) -> tuple[int, int, int]:
    year = max(1, number // 366 + 1)
    while _jalali_day_number(year + 1, 1, 1) <= number:
        year += 1
    while _jalali_day_number(year, 1, 1) > number:
        year -= 1
    remainder = number - _jalali_day_number(year, 1, 1)
    month = 1
    while True:
        if month <= 6:
            length = 31
        elif month <= 11:
            length = 30
        else:
            length = 30 if is_jalali_leap(year) else 29
        if remainder < length:
            return year, month, remainder + 1
        remainder -= length
        month += 1


def _jalali_date_string(day_number: int) -> str:
    year, month, day = _jalali_from_day_number(day_number)
    return f'{year:04d}/{month:02d}/{day:02d}'


def _jalali_weekday(date_str: str) -> str:
    try:
        year, month, day = parse_jalali_date(date_str)
    except ValueError:
        return ''
    today = get_today_jalali()
    today_index = (datetime.now().weekday() + 2) % 7
    today_number = _jalali_day_number(*today)
    return _WEEKDAY_NAMES[(today_index + _jalali_day_number(year, month, day) - today_number) % 7]


def _quick_range_bounds(range_key: str) -> tuple[str, str]:
    today = get_today_jalali()
    today_number = _jalali_day_number(*today)
    today_str = _jalali_date_string(today_number)
    if range_key == 'today':
        return today_str, today_str
    if range_key == 'week':
        return _jalali_date_string(today_number - 6), today_str
    if range_key == 'month':
        return f'{today[0]:04d}/{today[1]:02d}/01', today_str
    if range_key == 'last_month':
        if today[1] == 1:
            year, month = today[0] - 1, 12
        else:
            year, month = today[0], today[1] - 1
        first = _jalali_day_number(year, month, 1)
        if month <= 6:
            length = 31
        elif month <= 11:
            length = 30
        else:
            length = 30 if is_jalali_leap(year) else 29
        return _jalali_date_string(first), _jalali_date_string(first + length - 1)
    if range_key == 'term':
        start_year = today[0] if today[1] >= 7 else today[0] - 1
        return f'{start_year:04d}/07/01', today_str
    return '', ''


def _month_label(key: str) -> str:
    parts = (key or '').split('/')
    if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit() and 1 <= int(parts[1]) <= 12:
        return f'{MONTH_NAMES[int(parts[1]) - 1]} {parts[0]}'
    return key or 'بدون تاریخ'


def _report_filters(source=None) -> tuple[dict[str, object], str | None]:
    source = source if source is not None else request.args
    date_error = None
    raw_from = (source.get('date_from') or '').strip()[:10]
    raw_to = (source.get('date_to') or '').strip()[:10]
    range_key = (source.get('range') or '').strip()
    normalized_from = normalized_to = ''
    if raw_from or raw_to:
        try:
            normalized_from = normalize_jalali_date(raw_from) if raw_from else ''
            normalized_to = normalize_jalali_date(raw_to) if raw_to else ''
            if normalized_from and normalized_to and normalized_from > normalized_to:
                raise ValueError('بازه تاریخ نامعتبر است: تاریخ شروع بعد از تاریخ پایان است.')
        except ValueError as exc:
            date_error = str(exc)
            normalized_from, normalized_to = raw_from, raw_to

    if range_key in _RANGE_LABELS and range_key != 'custom':
        date_from, date_to = _quick_range_bounds(range_key)
    elif range_key == 'custom':
        date_from, date_to = normalized_from, normalized_to
    elif normalized_from or normalized_to:
        range_key = 'custom'
        date_from, date_to = normalized_from, normalized_to
    else:
        range_key = 'month'
        date_from, date_to = _quick_range_bounds('month')

    statuses = [status for status in source.getlist('status') if status in _STATUSES]
    sort = source.get('sort', 'date')
    if sort not in _SORT_COLUMNS:
        sort = 'date'
    direction = 'asc' if source.get('dir') == 'asc' else 'desc'
    try:
        per_page = int(source.get('per_page', 50))
    except (TypeError, ValueError):
        per_page = 50
    if per_page not in _PER_PAGE_OPTIONS:
        per_page = 50
    try:
        page = max(1, int(source.get('page', 1)))
    except (TypeError, ValueError):
        page = 1

    return {
        'range': range_key,
        'date_from': date_from,
        'date_to': date_to,
        'q': (source.get('q') or '').strip()[:80],
        'grade': (source.get('grade') or '').strip()[:60],
        'class_name': (source.get('class_name') or '').strip()[:60],
        'sida_class': (source.get('sida_class') or '').strip()[:60],
        'gender': (source.get('gender') or '').strip()[:20],
        'teacher_code': (source.get('teacher_code') or '').strip()[:40],
        'statuses': statuses,
        'sort': sort,
        'dir': direction,
        'page': page,
        'per_page': per_page,
    }, date_error


def _report_where(filters: dict[str, object], include_status: bool = True) -> tuple[str, list[object]]:
    clauses = ["(s.status='فعال' OR s.status IS NULL)"]
    params: list[object] = []
    if filters['date_from']:
        clauses.append('a.date>=?')
        params.append(filters['date_from'])
    if filters['date_to']:
        clauses.append('a.date<=?')
        params.append(filters['date_to'])
    if filters['q']:
        like = f"%{filters['q']}%"
        clauses.append('(s.first_name LIKE ? OR s.last_name LIKE ? OR s.code LIKE ?)')
        params.extend([like, like, like])
    for key, column in (('grade', 's.grade'), ('class_name', 's.class_name'), ('sida_class', 's.sida_class'),
                        ('gender', 's.gender'), ('teacher_code', 's.teacher_code')):
        if filters[key]:
            clauses.append(f'{column}=?')
            params.append(filters[key])
    if include_status and filters['statuses']:
        placeholders = ','.join('?' for _ in filters['statuses'])
        clauses.append(f'a.status IN ({placeholders})')
        params.extend(filters['statuses'])
    return ' AND '.join(clauses), params


def _report_analytics(rows) -> dict[str, object]:
    totals = {status: 0 for status in _STATUSES}
    students: dict[str, dict[str, object]] = {}
    months: dict[str, dict[str, int]] = {}
    weekdays = {name: {'absent': 0, 'total': 0} for name in _WEEKDAY_NAMES}
    classes: dict[str, dict[str, object]] = {}
    absent_days: dict[str, list[int]] = {}

    for row in rows:
        status = row['status']
        if status in totals:
            totals[status] += 1
        code = row['student_code']
        student = students.get(code)
        if student is None:
            student = {
                'code': code,
                'first_name': row['first_name'],
                'last_name': row['last_name'],
                'name': f"{row['first_name'] or ''} {row['last_name'] or ''}".strip(),
                'class_name': row['class_name'] or 'ثبت نشده',
                'grade': row['grade'] or 'ثبت نشده',
                'present': 0, 'absent': 0, 'excused': 0, 'late': 0, 'total': 0,
                'streak': 0, 'rate': None,
            }
            students[code] = student
            absent_days[code] = []
        student['total'] += 1
        counter_key = _STATUS_COUNTER_KEYS.get(status)
        if counter_key:
            student[counter_key] += 1

        date_value = (row['date'] or '').strip()
        if date_value:
            month = months.setdefault(date_value[:7], {'total': 0, 'present': 0, 'absent': 0, 'excused': 0, 'late': 0})
            month['total'] += 1
            if counter_key:
                month[counter_key] += 1
            weekday = _jalali_weekday(date_value)
            if weekday:
                weekdays[weekday]['total'] += 1
                if status == 'غایب':
                    weekdays[weekday]['absent'] += 1
            if status == 'غایب':
                try:
                    absent_days[code].append(_jalali_day_number(*parse_jalali_date(date_value)))
                except ValueError:
                    pass

        class_name = row['class_name'] or 'ثبت نشده'
        klass = classes.setdefault(class_name, {'total': 0, 'absent': 0, 'late': 0, 'student_codes': set()})
        klass['total'] += 1
        klass['absent'] += 1 if status == 'غایب' else 0
        klass['late'] += 1 if status == 'تاخیر' else 0
        klass['student_codes'].add(code)

    for code, student in students.items():
        total = student['total'] or 0
        student['rate'] = round(student['present'] * 100 / total) if total else None
        best = current = 0
        previous = None
        for day in sorted(set(absent_days[code])):
            current = current + 1 if previous is not None and day - previous == 1 else 1
            best = max(best, current)
            previous = day
        student['streak'] = best

    total_marks = sum(totals.values())
    month_trend = [{'key': key, 'label': _month_label(key), **months[key]} for key in sorted(months)]
    classes_summary = sorted(
        ({'class_name': name, 'students': len(data['student_codes']), 'total': data['total'],
          'absent': data['absent'], 'late': data['late'],
          'rate': round((data['total'] - data['absent']) * 100 / data['total']) if data['total'] else None}
         for name, data in classes.items()),
        key=lambda item: (-item['absent'], item['class_name']),
    )
    return {
        'totals': totals,
        'total_marks': total_marks,
        'attendance_rate': round(totals['حاضر'] * 100 / total_marks) if total_marks else None,
        'absent_students': sum(1 for data in students.values() if data['absent']),
        'late_students': sum(1 for data in students.values() if data['late']),
        'students': sorted(students.values(), key=lambda item: (-item['absent'], -item['late'], item['name'])),
        'month_trend': month_trend,
        'weekday_trend': [{'label': name, **weekdays[name]} for name in _WEEKDAY_NAMES],
        'classes': classes_summary[:15],
        'class_count': len(classes_summary),
        'max_month_total': max((month['total'] for month in months.values()), default=0),
        'max_weekday_total': max((day['total'] for day in weekdays.values()), default=0),
        'max_class_absent': max((item['absent'] for item in classes_summary), default=0),
    }


def _report_query(**overrides) -> dict[str, object]:
    query: dict[str, object] = {}
    for key in ('range', 'date_from', 'date_to', 'q', 'grade', 'class_name', 'sida_class', 'gender',
                'teacher_code', 'sort', 'dir', 'per_page', 'print', 'autoprint', 'print_all'):
        value = request.args.get(key)
        if value:
            query[key] = value
    statuses = [status for status in request.args.getlist('status') if status in _STATUSES]
    if statuses:
        query['status'] = statuses
    for key, value in overrides.items():
        if value in (None, '', []):
            query.pop(key, None)
        else:
            query[key] = value
    return query


def _filters_description(filters: dict[str, object], teachers_list) -> str:
    parts: list[str] = []
    if filters['date_from'] and filters['date_to']:
        parts.append(f"از {filters['date_from']} تا {filters['date_to']}")
    elif filters['date_from']:
        parts.append(f"از {filters['date_from']} به بعد")
    elif filters['date_to']:
        parts.append(f"تا {filters['date_to']}")
    else:
        parts.append('همه بازه')
    if filters['q']:
        parts.append(f"جست‌وجو: {filters['q']}")
    for key, label in (('grade', 'پایه'), ('class_name', 'کلاس'), ('sida_class', 'سیدا'), ('gender', 'جنسیت')):
        if filters[key]:
            parts.append(f'{label}: {filters[key]}')
    if filters['teacher_code']:
        name = next((f"{teacher['first_name']} {teacher['last_name']}"
                     for teacher in teachers_list if teacher['code'] == filters['teacher_code']),
                    filters['teacher_code'])
        parts.append(f'معلم: {name}')
    if filters['statuses']:
        parts.append('وضعیت: ' + '، '.join(filters['statuses']))
    return ' | '.join(parts)


def _excel_value(value):
    """Return a value safe for spreadsheet display, preventing formula injection."""
    if value is None:
        return ''
    if isinstance(value, str) and value[:1] in {'=', '+', '-', '@'}:
        return "'" + value
    return value


def _redirect_with_filters(endpoint):
    target = request.form.get('next', '').strip()
    return redirect(target or url_for(endpoint))


def _save_driver_destinations(conn, driver_id):
    previous_assignments = conn.execute(
        '''SELECT s.id, dest.destination
           FROM students s JOIN service_driver_destinations dest
             ON dest.id=s.service_destination_id
           WHERE s.service_driver_id=?''',
        (driver_id,),
    ).fetchall()
    values = []
    for raw in request.form.getlist('destinations[]'):
        destination = raw.strip()[:120]
        if destination and destination not in values:
            values.append(destination)
    conn.execute('DELETE FROM service_driver_destinations WHERE driver_id=?', (driver_id,))
    conn.executemany(
        'INSERT INTO service_driver_destinations(driver_id,destination) VALUES(?,?)',
        [(driver_id, destination) for destination in values[:20]],
    )
    destination_ids = {
        row['destination']: row['id']
        for row in conn.execute(
            'SELECT id,destination FROM service_driver_destinations WHERE driver_id=?',
            (driver_id,),
        ).fetchall()
    }
    for assignment in previous_assignments:
        conn.execute(
            'UPDATE students SET service_destination_id=? WHERE id=?',
            (destination_ids.get(assignment['destination']), assignment['id']),
        )


def _service_destination_id(conn, driver_id, raw_destination_id):
    if not driver_id or not raw_destination_id or not str(raw_destination_id).isdigit():
        return None
    row = conn.execute(
        'SELECT id FROM service_driver_destinations WHERE id=? AND driver_id=?',
        (int(raw_destination_id), int(driver_id)),
    ).fetchone()
    return row['id'] if row else None


def service_drivers():
    with get_db() as conn:
        if request.method == 'POST':
            action = request.form.get('action', '').strip()
            backed_up = False
            if action == 'save_driver':
                driver_id = request.form.get('driver_id', '').strip()
                driver_name = request.form.get('driver_name', '').strip()[:120]
                phone = request.form.get('phone', '').strip()[:40]
                vehicle = request.form.get('vehicle', '').strip()[:80]
                plate = request.form.get('plate', '').strip()[:30]
                if not driver_name:
                    flash('نام راننده الزامی است.', 'danger')
                else:
                    auto_backup()
                    backed_up = True
                    if driver_id.isdigit() and conn.execute(
                            'SELECT 1 FROM service_driver_registry WHERE id=?', (driver_id,)).fetchone():
                        conn.execute('''UPDATE service_driver_registry
                                        SET driver_name=?,phone=?,vehicle=?,plate=? WHERE id=?''',
                                      (driver_name, phone, vehicle, plate, driver_id))
                        saved_driver_id = int(driver_id)
                    else:
                        cursor = conn.execute('''INSERT INTO service_driver_registry(driver_name,phone,vehicle,plate)
                                                VALUES(?,?,?,?)''', (driver_name, phone, vehicle, plate))
                        saved_driver_id = cursor.lastrowid
                    _save_driver_destinations(conn, saved_driver_id)
                    conn.commit()
                    audit(session.get('user_id'), 'save_service_driver', 'service_driver', str(saved_driver_id), driver_name)
                    flash('اطلاعات راننده ذخیره شد.', 'success')
            elif action == 'delete_driver':
                driver_id = request.form.get('driver_id', '').strip()
                if driver_id.isdigit():
                    auto_backup()
                    conn.execute('''UPDATE students
                                    SET service_driver_id=NULL, service_destination_id=NULL
                                    WHERE service_driver_id=?''', (driver_id,))
                    conn.execute('DELETE FROM service_driver_destinations WHERE driver_id=?', (driver_id,))
                    conn.execute('DELETE FROM service_driver_registry WHERE id=?', (driver_id,))
                    conn.commit()
                    audit(session.get('user_id'), 'delete_service_driver', 'service_driver', driver_id)
                    flash('راننده حذف شد و انتساب دانش‌آموزان آن برداشته شد.', 'info')
            elif action == 'assign_student':
                student_id = request.form.get('student_id', '').strip()
                driver_id = request.form.get('driver_id', '').strip()
                destination_id = None
                if driver_id and not driver_id.isdigit():
                    driver_id = ''
                if driver_id and not conn.execute(
                        'SELECT 1 FROM service_driver_registry WHERE id=?', (driver_id,)).fetchone():
                    driver_id = ''
                if driver_id:
                    destination_id = _service_destination_id(
                        conn, driver_id, request.form.get('destination_id', '').strip()
                    )
                    raw_destination_id = request.form.get('destination_id', '').strip()
                    if raw_destination_id and destination_id is None:
                        flash('مقصد انتخاب‌شده برای این راننده معتبر نیست.', 'danger')
                        return redirect(url_for('service_drivers'))
                if student_id.isdigit() and conn.execute(
                        "SELECT 1 FROM students WHERE id=? AND (status='فعال' OR status IS NULL)",
                        (student_id,),
                ).fetchone():
                    auto_backup()
                    conn.execute(
                        'UPDATE students SET service_driver_id=?, service_destination_id=? WHERE id=?',
                        (int(driver_id) if driver_id else None, destination_id, int(student_id)),
                    )
                    conn.commit()
                    audit(session.get('user_id'), 'assign_service_student', 'student', student_id, driver_id or 'none')
                    flash('راننده و مقصد دانش‌آموز ذخیره شد.', 'success')
            elif action == 'assign_students':
                student_ids = [value for value in request.form.getlist('student_ids') if value.isdigit()]
                driver_id = request.form.get('driver_id', '').strip()
                raw_destination_id = request.form.get('destination_id', '').strip()
                if driver_id and not driver_id.isdigit():
                    driver_id = ''
                if driver_id and not conn.execute(
                        'SELECT 1 FROM service_driver_registry WHERE id=?', (driver_id,)).fetchone():
                    driver_id = ''
                destination_id = _service_destination_id(conn, driver_id, raw_destination_id) if driver_id else None
                if raw_destination_id and destination_id is None:
                    flash('مقصد انتخاب‌شده برای این راننده معتبر نیست.', 'danger')
                    return redirect(url_for('service_drivers'))
                if not student_ids:
                    flash('حداقل یک دانش‌آموز را انتخاب کنید.', 'warning')
                else:
                    auto_backup()
                    placeholders = ','.join('?' for _ in student_ids)
                    conn.execute(
                        f'''UPDATE students SET service_driver_id=?, service_destination_id=?
                            WHERE id IN ({placeholders}) AND (status='فعال' OR status IS NULL)''',
                        [int(driver_id) if driver_id else None, destination_id, *student_ids],
                    )
                    conn.commit()
                    audit(session.get('user_id'), 'assign_service_driver', 'student', ','.join(student_ids), driver_id or 'none')
                    flash(f'راننده و مقصد برای {len(student_ids)} دانش‌آموز ثبت شد.', 'success')
            elif action == 'save_all_destinations':
                import json as _json
                try:
                    data = _json.loads(request.form.get('destinations_json', '{}'))
                except (TypeError, ValueError):
                    data = {}
                if not isinstance(data, dict):
                    data = {}
                pending = []
                for driver_id_str, destinations in data.items():
                    if not driver_id_str.isdigit():
                        continue
                    driver_id = int(driver_id_str)
                    if not conn.execute(
                        'SELECT 1 FROM service_driver_registry WHERE id=?', (driver_id,)).fetchone():
                        continue
                    if not isinstance(destinations, list):
                        destinations = []
                    values = []
                    for raw in destinations:
                        destination = str(raw).strip()[:120]
                        if destination and destination not in values:
                            values.append(destination)
                    pending.append((driver_id, values[:20]))

                if pending:
                    auto_backup()
                for driver_id, values in pending:
                    previous_assignments = conn.execute(
                        '''SELECT s.id, dest.destination
                           FROM students s JOIN service_driver_destinations dest
                             ON dest.id=s.service_destination_id
                           WHERE s.service_driver_id=?''',
                        (driver_id,),
                    ).fetchall()
                    conn.execute('DELETE FROM service_driver_destinations WHERE driver_id=?', (driver_id,))
                    conn.executemany(
                        'INSERT INTO service_driver_destinations(driver_id,destination) VALUES(?,?)',
                        [(driver_id, destination) for destination in values],
                    )
                    destination_ids = {
                        row['destination']: row['id']
                        for row in conn.execute(
                            'SELECT id,destination FROM service_driver_destinations WHERE driver_id=?',
                            (driver_id,),
                        ).fetchall()
                    }
                    for assignment in previous_assignments:
                        conn.execute(
                            'UPDATE students SET service_destination_id=? WHERE id=?',
                            (destination_ids.get(assignment['destination']), assignment['id']),
                        )
                saved_count = len(pending)
                if saved_count:
                    conn.commit()
                    audit(session.get('user_id'), 'save_all_destinations', 'service_driver', str(saved_count), 'bulk')
                    flash(f'مقصدهای {saved_count} راننده ذخیره شد.', 'success')
                else:
                    flash('مقصد معتبری برای ذخیره پیدا نشد.', 'warning')
            return redirect(url_for('service_drivers'))

        driver_rows = conn.execute('''SELECT d.id,d.driver_name,d.phone,d.vehicle,d.plate,
                                          COUNT(CASE WHEN s.status='فعال' OR s.status IS NULL THEN s.id END) AS student_count
                                  FROM service_driver_registry d
                                  LEFT JOIN students s ON s.service_driver_id=d.id
                                  GROUP BY d.id
                                  ORDER BY d.driver_name''').fetchall()
        drivers = []
        for row in driver_rows:
            driver = dict(row)
            driver['destinations'] = [dict(item) for item in conn.execute(
                'SELECT id,destination FROM service_driver_destinations WHERE driver_id=? ORDER BY id',
                (driver['id'],),
            ).fetchall()]
            drivers.append(driver)
        edit_driver = None
        edit_id = request.args.get('edit', '').strip()
        if edit_id.isdigit():
            row = conn.execute('SELECT id,driver_name,phone,vehicle,plate FROM service_driver_registry WHERE id=?', (edit_id,)).fetchone()
            if row:
                edit_driver = dict(row)
                edit_driver['destinations'] = [dict(item) for item in conn.execute(
                    'SELECT id,destination FROM service_driver_destinations WHERE driver_id=? ORDER BY id',
                    (edit_id,),
                ).fetchall()]
        students = conn.execute('''SELECT s.id,s.first_name,s.last_name,s.code,s.grade,s.class_name,
                                           s.service_driver_id,s.service_destination_id,
                                           d.driver_name, dest.destination AS service_destination,
                                           TRIM(COALESCE(t.first_name, '') || ' ' || COALESCE(t.last_name, '')) AS teacher_name
                                   FROM students s
                                   LEFT JOIN service_driver_registry d ON d.id=s.service_driver_id
                                   LEFT JOIN service_driver_destinations dest ON dest.id=s.service_destination_id
                                   LEFT JOIN teachers t ON TRIM(t.code)=TRIM(s.teacher_code)
                                   WHERE s.status='فعال' OR s.status IS NULL
                                   ORDER BY s.class_name,s.first_name,s.last_name''').fetchall()
        destination_map = {
            str(driver['id']): [
                {'id': destination['id'], 'destination': destination['destination']}
                for destination in driver['destinations']
            ]
            for driver in drivers
        }
    return render_template(
        'service_drivers.html', drivers=drivers, students=students,
        edit_driver=edit_driver, destination_map=destination_map,
    )


def service_drivers_print():
    with get_db() as conn:
        destination_rows = conn.execute('''SELECT dest.destination, d.id AS driver_id, d.driver_name, d.phone, d.vehicle, d.plate
                                           FROM service_driver_destinations dest
                                           JOIN service_driver_registry d ON d.id=dest.driver_id
                                           ORDER BY dest.destination, d.driver_name''').fetchall()
        student_rows = conn.execute('''SELECT COALESCE(dest.destination, 'بدون مقصد مشخص') AS destination,
                                               d.id AS driver_id, d.driver_name, d.phone, d.vehicle, d.plate,
                                               s.id AS student_id, s.first_name, s.last_name, s.photo, s.class_name, s.grade,
                                               TRIM(COALESCE(t.first_name, '') || ' ' || COALESCE(t.last_name, '')) AS teacher_name
                                        FROM students s
                                        JOIN service_driver_registry d ON d.id=s.service_driver_id
                                        LEFT JOIN service_driver_destinations dest
                                          ON dest.id=s.service_destination_id AND dest.driver_id=d.id
                                        LEFT JOIN teachers t ON TRIM(t.code)=TRIM(s.teacher_code)
                                        WHERE s.status='فعال' OR s.status IS NULL
                                        ORDER BY destination, d.driver_name, s.last_name, s.first_name''').fetchall()
        driver_layout = prepare_for_print(load_layout(conn, 'service_drivers', session.get('user_id')))

    groups = {}

    def add_row(destination, row):
        group = groups.setdefault(destination, {'destination': destination, 'drivers': [], 'student_count': 0})
        driver = next((item for item in group['drivers'] if item['id'] == row['driver_id']), None)
        if driver is None:
            driver = {
                'id': row['driver_id'], 'driver_name': row['driver_name'], 'phone': row['phone'],
                'vehicle': row['vehicle'], 'plate': row['plate'], 'students': [],
            }
            group['drivers'].append(driver)
        student_id = row['student_id'] if 'student_id' in row.keys() else None
        if student_id:
            driver['students'].append({
                'id': student_id, 'first_name': row['first_name'], 'last_name': row['last_name'],
                'photo': row['photo'], 'class_name': row['class_name'], 'grade': row['grade'],
                'teacher_name': row['teacher_name'],
            })
            group['student_count'] += 1

    for row in destination_rows:
        add_row(row['destination'], row)
    for row in student_rows:
        add_row((row['destination'] or 'بدون مقصد ثبت‌شده').strip(), row)

    destination_groups = sorted(groups.values(), key=lambda group: group['destination'])
    return render_template('service_drivers_print.html', destination_groups=destination_groups,
                           driver_layout=driver_layout)


def class_management():
    with get_db() as conn:
        if request.method == 'POST':
            student_ids = [sid.strip() for sid in request.form.getlist('student_ids') if sid.strip()]
            valid_teacher_codes = {row['code'] for row in conn.execute('SELECT code FROM teachers').fetchall()}
            updated_count = 0
            invalid_teacher_count = 0
            not_found_count = 0
            backed_up = False
            pending_audits = []
            for student_id in student_ids:
                row = conn.execute('SELECT class_name,sida_class,teacher_code FROM students WHERE id=?', (student_id,)).fetchone()
                if not row:
                    not_found_count += 1
                    continue
                class_name = request.form.get(f'class_name__{student_id}', '').strip()[:80]
                sida_class = request.form.get(f'sida_class__{student_id}', '').strip()[:80]
                teacher_code = request.form.get(f'teacher_code__{student_id}', '').strip()
                if teacher_code and teacher_code not in valid_teacher_codes:
                    invalid_teacher_count += 1
                    continue
                if (row['class_name'] or '') == class_name and (row['sida_class'] or '') == sida_class \
                        and (row['teacher_code'] or '') == teacher_code:
                    continue
                if not backed_up:
                    auto_backup()
                    backed_up = True
                conn.execute('UPDATE students SET class_name=?,sida_class=?,teacher_code=? WHERE id=?',
                             (class_name, sida_class, teacher_code, student_id))
                pending_audits.append((student_id, f'{class_name}|{sida_class}|{teacher_code}'))
                updated_count += 1
            conn.commit()
            for student_id, details in pending_audits:
                audit(session.get('user_id'), 'update_class_assignment', 'student', student_id, details)
            if updated_count:
                flash(f'✅ اطلاعات {updated_count} دانش‌آموز با موفقیت ذخیره شد.', 'success')
            if invalid_teacher_count:
                flash(f'⚠️ {invalid_teacher_count} ردیف به‌دلیل معلم نامعتبر ذخیره نشد.', 'danger')
            if not_found_count:
                flash(f'⚠️ {not_found_count} ردیف یافت نشد و ذخیره نشد.', 'danger')
            if not updated_count and not invalid_teacher_count and not not_found_count:
                flash('تغییری برای ذخیره یافت نشد.', 'info')
            return redirect(url_for('class_management'))
        students = conn.execute('''SELECT s.id,s.first_name,s.last_name,s.code,s.class_name,s.sida_class,s.teacher_code,
                                          t.first_name AS teacher_first,t.last_name AS teacher_last
                                   FROM students s LEFT JOIN teachers t ON s.teacher_code=t.code
                                   WHERE s.status='فعال' OR s.status IS NULL
                                   ORDER BY s.class_name,s.first_name,s.last_name''').fetchall()
        teachers = conn.execute('SELECT code,first_name,last_name,class_name,sida_class FROM teachers ORDER BY first_name,last_name').fetchall()
    return render_template('class_management.html', students=students, teachers=teachers,
                           sida_class_options=SIDA_CLASS_OPTIONS)

def _fa_digits(value) -> str:
    """Render a number with Persian digits for flash messages."""
    return str(value).translate(str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹'))


def _quick_scope(class_filter: str) -> tuple[str, list[object]]:
    """Build the active-student scope used by the quick attendance list."""
    scope_sql = "(s.status='فعال' OR s.status IS NULL)"
    scope_params: list[object] = []
    if class_filter:
        scope_sql += ' AND s.class_name = ?'
        scope_params.append(class_filter)
    return scope_sql, scope_params


def _attendance_quick():
    """ثبت سریع حضور و غیاب کل مدرسه با انتخاب فقط غایبین.

    مدیر فقط اسامی غایبین را انتخاب میکند؛ همهٔ دانشآموزان فعال دیگر که برای
    تاریخ انتخابشده رکورد ندارند بهصورت خودکار «حاضر» ثبت میشوند. رکوردهای
    موجود (مرخصی، تاخیر و ...) هرگز بازنویسی نمیشوند.
    """
    with get_db() as conn:
        if request.method == 'POST':
            try:
                date_value = normalize_jalali_date(request.form.get('date', ''), required=True)
            except ValueError as exc:
                flash(f'⚠️ {exc}', 'danger')
                return redirect(url_for('attendance_quick'))
            class_filter = (request.form.get('class_name', '').strip() or '')[:80]
            absent_codes = {
                key[len('absent_'):]
                for key, value in request.form.items()
                if key.startswith('absent_') and value == '1'
            }
            scope_sql, scope_params = _quick_scope(class_filter)
            eligible = {
                row['code'] for row in conn.execute(
                    f'''SELECT s.code FROM students s
                        WHERE {scope_sql}
                          AND NOT EXISTS (SELECT 1 FROM attendance_students a
                                          WHERE a.student_code = s.code AND a.date = ?)
                        ORDER BY s.class_name, s.last_name, s.first_name''',
                    [*scope_params, date_value],
                ).fetchall()
            }
            absent_codes &= eligible
            if eligible:
                auto_backup()
                conn.executemany(
                    '''INSERT INTO attendance_students(student_code,date,status) VALUES(?,?,?)
                       ON CONFLICT(student_code,date) DO NOTHING''',
                    [(code, date_value, 'غایب' if code in absent_codes else 'حاضر')
                     for code in sorted(eligible)],
                )
                conn.commit()
                present_count = len(eligible) - len(absent_codes)
                audit(session.get('user_id'), 'quick_student_attendance', 'attendance_students',
                      date_value,
                      f'حاضر {present_count} | غایب {len(absent_codes)} | کل {len(eligible)}'
                      + (f' | کلاس {class_filter}' if class_filter else ''))
                flash(
                    f'✅ حضور و غیاب {_fa_digits(date_value)} ثبت شد: '
                    f'{_fa_digits(present_count)} حاضر، {_fa_digits(len(absent_codes))} غایب '
                    f'(رکوردهای موجود تغییر نکرد).',
                    'success',
                )
            else:
                flash('دانشآموز فعالی برای ثبت در این بازه یافت نشد.', 'warning')
            return redirect(url_for('attendance_quick', date=date_value, class_name=class_filter))

        date_value = (request.args.get('date', '').strip() or today_string())
        try:
            date_value = normalize_jalali_date(date_value)
        except ValueError:
            date_value = today_string()
        class_filter = (request.args.get('class_name', '').strip() or '')[:80]
        scope_sql, scope_params = _quick_scope(class_filter)
        students = [dict(row) for row in conn.execute(
            f'''SELECT s.id, s.first_name, s.last_name, s.code, s.grade, s.class_name, s.sida_class,
                       a.status AS current_status,
                       (SELECT COUNT(*) FROM attendance_students ax
                         WHERE ax.student_code = s.code AND ax.status = 'غایب') AS total_absences
                FROM students s
                LEFT JOIN attendance_students a ON a.student_code = s.code AND a.date = ?
                WHERE {scope_sql}
                ORDER BY s.class_name, s.last_name, s.first_name''',
            [date_value, *scope_params],
        ).fetchall()]
        classes = conn.execute(
            "SELECT DISTINCT class_name FROM students WHERE class_name IS NOT NULL AND class_name != '' "
            "AND (status='فعال' OR status IS NULL) ORDER BY class_name"
        ).fetchall()
        marked_dates = [
            row['date'] for row in conn.execute(
                'SELECT DISTINCT date FROM attendance_students WHERE date LIKE ?',
                (f'{date_value[:7]}/%',),
            ).fetchall()
        ]
        already_marked = conn.execute(
            'SELECT COUNT(*) AS count FROM attendance_students WHERE date = ?', (date_value,)
        ).fetchone()['count']

    jy, jm, jd = get_today_jalali()
    return render_template(
        'attendance_students.html',
        quick_mode=True,
        students=students,
        classes=classes,
        class_filter=class_filter,
        date_value=date_value,
        today_value=today_string(),
        yesterday_value=_jalali_date_string(_jalali_day_number(jy, jm, jd) - 1),
        today_jalali={'year': jy, 'month': jm, 'day': jd},
        today_weekday_index=(datetime.now().weekday() + 2) % 7,
        date_weekday=_jalali_weekday(date_value),
        years=list(range(max(1, jy - 150), jy + 26)),
        marked_dates=marked_dates,
        already_marked=already_marked,
        statuses=_STATUSES,
    )


def attendance_students():
    scope = session.get('personnel_number') if session.get('role') == 'teacher' else None
    with get_db() as conn:
        requested_teacher = (scope or request.args.get('teacher_code', '').strip()
                             or request.form.get('teacher_code', '').strip())
        if request.method == 'POST':
            try:
                date = normalize_jalali_date(request.form.get('date', ''), required=True)
            except ValueError as exc:
                flash(f'⚠️ {exc}', 'danger')
                return redirect(url_for('attendance_students'))
            if not requested_teacher:
                flash('⚠️ ابتدا یک معلم انتخاب کنید.', 'danger')
                return redirect(url_for('attendance_students'))

            values = {}
            for key, value in request.form.items():
                if key.startswith('status_') and value in _STATUSES:
                    values[key.removeprefix('status_')] = value

            allowed = {
                row['code'] for row in conn.execute(
                    'SELECT code FROM students WHERE (status="فعال" OR status IS NULL) AND teacher_code=?',
                    (requested_teacher,),
                ).fetchall()
            }
            entries = [(code, date, status) for code, status in values.items() if code in allowed]
            if entries:
                auto_backup()
                conn.executemany(
                    '''INSERT INTO attendance_students(student_code,date,status) VALUES(?,?,?)
                       ON CONFLICT(student_code,date) DO UPDATE SET status=excluded.status''',
                    entries,
                )
                conn.commit()
                counts = {}
                for _, _, status in entries:
                    counts[status] = counts.get(status, 0) + 1
                summary = '، '.join(f'{status} {count} نفر' for status, count in counts.items())
                audit(session.get('user_id'), 'save_student_attendance', 'attendance_students',
                      f'{requested_teacher}/{date}', summary)
                flash(f'✅ حضور و غیاب {len(entries)} دانش‌آموز ثبت/به‌روزرسانی شد ({summary}).', 'success')
            else:
                flash('رکورد معتبری برای ذخیره پیدا نشد.', 'warning')
            target = {'date': date}
            if requested_teacher:
                target['teacher_code'] = requested_teacher
            return redirect(url_for('attendance_students', **target))

        if scope:
            teachers = conn.execute(
                'SELECT id, first_name, last_name, code, class_name FROM teachers WHERE code=? ORDER BY first_name',
                (scope,),
            ).fetchall()
        else:
            teachers = conn.execute('SELECT id, first_name, last_name, code, class_name FROM teachers ORDER BY first_name').fetchall()

        date_value = request.args.get('date', '').strip()
        try:
            date_value = normalize_jalali_date(date_value) if date_value else today_string()
        except ValueError:
            date_value = today_string()

        selected_teacher = None
        students = []
        date_stats = {}
        marked_dates = []
        if requested_teacher:
            selected_teacher = conn.execute('SELECT * FROM teachers WHERE code=?', (requested_teacher,)).fetchone()
            if selected_teacher and (scope is None or requested_teacher == scope):
                year, month, day = (int(part) for part in date_value.split('/'))
                recent_from = _jalali_date_string(max(1, _jalali_day_number(year, month, day) - 30))
                rows = conn.execute(
                    '''SELECT s.id, s.first_name, s.last_name, s.code, s.grade, s.class_name, s.sida_class,
                              a.status AS current_status,
                              (SELECT COUNT(*) FROM attendance_students ax
                               WHERE ax.student_code = s.code AND ax.status='غایب'
                                 AND ax.date >= ? AND ax.date <= ?) AS recent_absences
                       FROM students s
                       LEFT JOIN attendance_students a ON a.student_code = s.code AND a.date = ?
                       WHERE (s.status='فعال' OR s.status IS NULL) AND s.teacher_code = ?
                       ORDER BY s.last_name, s.first_name''',
                    (recent_from, date_value, date_value, requested_teacher),
                ).fetchall()
                students = [dict(row) for row in rows]
                for row in conn.execute(
                    '''SELECT a.status, COUNT(*) AS count
                       FROM attendance_students a JOIN students s ON s.code = a.student_code
                       WHERE s.teacher_code = ? AND a.date = ?
                       GROUP BY a.status''',
                    (requested_teacher, date_value),
                ).fetchall():
                    date_stats[row['status']] = row['count']
                month_prefix = f'{date_value[:7]}/%'
                marked_dates = [
                    row['date'] for row in conn.execute(
                        '''SELECT DISTINCT a.date FROM attendance_students a
                           JOIN students s ON s.code = a.student_code
                           WHERE s.teacher_code = ? AND a.date LIKE ?''',
                        (requested_teacher, month_prefix),
                    ).fetchall()
                ]

        jy, jm, jd = get_today_jalali()
        today_jalali = {'year': jy, 'month': jm, 'day': jd}
        today_value = today_string()
        yesterday_value = _jalali_date_string(_jalali_day_number(jy, jm, jd) - 1)
        return render_template(
            'attendance_students.html',
            teachers=teachers,
            students=students,
            selected_teacher=selected_teacher,
            teacher_code=requested_teacher,
            years=list(range(max(1, jy - 150), jy + 26)),
            today_jalali=today_jalali,
            date_value=date_value,
            today_value=today_value,
            yesterday_value=yesterday_value,
            today_weekday_index=(datetime.now().weekday() + 2) % 7,
            date_weekday=_jalali_weekday(date_value),
            date_stats=date_stats,
            marked_dates=marked_dates,
            statuses=_STATUSES,
            recent_absent_students=sum(1 for student in students if student['recent_absences']),
        )


def attendance_teachers():
    with get_db() as conn:
        if request.method == 'POST':
            try:
                date = normalize_jalali_date(request.form.get('date', ''), required=True)
            except ValueError as exc:
                flash(f'⚠️ {exc}', 'danger')
                return redirect(url_for('attendance_teachers'))
            allowed_codes = {
                row['code'] for row in conn.execute('SELECT code FROM teachers').fetchall()
            }
            pending = []
            for key, value in request.form.items():
                if not key.startswith('status_') or value not in _STATUSES:
                    continue
                code = key.removeprefix('status_')
                if code not in allowed_codes:
                    continue
                try:
                    late = max(0, int(request.form.get(f'late_duration_{code}', '0') or 0))
                except ValueError:
                    late = 0
                pending.append((code, value, late))

            if pending:
                auto_backup()
            for code, value, late in pending:
                existing = conn.execute('SELECT id FROM attendance_teachers WHERE teacher_code=? AND date=?', (code, date)).fetchone()
                if existing:
                    conn.execute('UPDATE attendance_teachers SET status=?, late_duration=? WHERE id=?', (value, late, existing['id']))
                else:
                    conn.execute('INSERT INTO attendance_teachers(teacher_code,date,status,late_duration) VALUES(?,?,?,?)', (code, date, value, late))
            if pending:
                conn.commit()
            saved_count = len(pending)
            flash(f'✅ حضور و غیاب {saved_count} معلم با موفقیت ثبت شد.', 'success')
            return redirect(url_for('attendance_teachers', date=date))

        jy, jm, jd = get_today_jalali()
        years = list(range(max(1, jy - 150), jy + 26))
        today_jalali = {'year': jy, 'month': jm, 'day': jd}
        date_value = request.args.get('date', '').strip()
        try:
            date_value = normalize_jalali_date(date_value) if date_value else today_string()
        except ValueError:
            date_value = today_string()

        teachers = conn.execute(
            '''SELECT t.*, a.status AS current_status,
                      COALESCE(a.late_duration, 0) AS current_late_duration
               FROM teachers t
               LEFT JOIN attendance_teachers a
                 ON a.teacher_code = t.code AND a.date = ?
               ORDER BY t.first_name, t.last_name''',
            (date_value,),
        ).fetchall()
        date_stats = {
            row['status']: row['count']
            for row in conn.execute(
                '''SELECT status, COUNT(*) AS count
                   FROM attendance_teachers WHERE date=? GROUP BY status''',
                (date_value,),
            ).fetchall()
        }
        month_prefix = f'{date_value[:7]}/%'
        marked_dates = [
            row['date'] for row in conn.execute(
                '''SELECT DISTINCT date FROM attendance_teachers
                   WHERE date LIKE ? ORDER BY date''',
                (month_prefix,),
            ).fetchall()
        ]
        today_value = today_string()
        yesterday_value = _jalali_date_string(_jalali_day_number(jy, jm, jd) - 1)
    return render_template(
        'attendance_teachers.html',
        teachers=teachers,
        years=years,
        today_jalali=today_jalali,
        today_value=today_value,
        yesterday_value=yesterday_value,
        date_value=date_value,
        date_weekday=_jalali_weekday(date_value),
        today_weekday_index=(datetime.now().weekday() + 2) % 7,
        date_stats=date_stats,
        marked_dates=marked_dates,
        statuses=_STATUSES,
    )


def delete_student_attendance(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM attendance_students WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_student_attendance', 'attendance_students', id)
        flash('✅ رکورد حضور و غیاب دانش‌آموز حذف شد.', 'success')
    else:
        flash('رکورد حضور و غیاب پیدا نشد.', 'warning')
    return _redirect_with_filters('attendance_report')


def update_student_attendance(id):
    status = request.form.get('status', '').strip()
    raw_date = request.form.get('date', '').strip()
    if status not in _STATUSES:
        flash('وضعیت حضور و غیاب نامعتبر است.', 'danger')
        return _redirect_with_filters('attendance_report')
    try:
        date = normalize_jalali_date(raw_date)
    except ValueError as exc:
        flash(f'⚠️ {exc}', 'danger')
        return _redirect_with_filters('attendance_report')
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('UPDATE attendance_students SET date=?,status=? WHERE id=?', (date, status, id))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'update_student_attendance', 'attendance_students', id, f'{date} | {status}')
        flash('✅ رکورد حضور و غیاب ویرایش شد.', 'success')
    else:
        flash('رکورد حضور و غیاب پیدا نشد.', 'warning')
    return _redirect_with_filters('attendance_report')


def delete_teacher_attendance(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM attendance_teachers WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_teacher_attendance', 'attendance_teachers', id)
        flash('✅ رکورد حضور و غیاب معلم حذف شد.', 'success')
    else:
        flash('رکورد حضور و غیاب پیدا نشد.', 'warning')
    return _redirect_with_filters('teacher_attendance_report')


def delete_filtered_student_absences():
    filters, date_error = _report_filters(request.form)
    if date_error:
        flash(f'⚠️ {date_error}', 'danger')
        return _redirect_with_filters('attendance_report')
    where, params = _report_where(filters, include_status=False)
    auto_backup()
    with get_db() as conn:
        ids = conn.execute(
            f'''SELECT a.id FROM attendance_students a JOIN students s ON a.student_code=s.code
                WHERE {where} AND a.status='غایب' ''',
            params,
        ).fetchall()
        if ids:
            marks = ','.join('?' for _ in ids)
            conn.execute(f'DELETE FROM attendance_students WHERE id IN ({marks})', [row['id'] for row in ids])
        conn.commit()
    audit(session.get('user_id'), 'delete_filtered_student_absences', 'attendance_students', '',
          f'{len(ids)} رکورد | {_filters_description(filters, [])}')
    if ids:
        flash(f'✅ {len(ids)} رکورد غیبت دانش‌آموز حذف شد.', 'success')
    else:
        flash('رکورد غیبتی برای حذف یافت نشد.', 'info')
    return _redirect_with_filters('attendance_report')


def delete_filtered_teacher_absences():
    date_from, date_to = request.form.get('date_from', '').strip(), request.form.get('date_to', '').strip()
    try:
        date_from = normalize_jalali_date(date_from) if date_from else ''
        date_to = normalize_jalali_date(date_to) if date_to else ''
    except ValueError as exc:
        flash(f'⚠️ {exc}', 'danger')
        return _redirect_with_filters('teacher_attendance_report')
    where = ['a.status="غایب"']
    params = []
    if date_from: where.append('a.date>=?'); params.append(date_from)
    if date_to: where.append('a.date<=?'); params.append(date_to)
    if request.form.get('teacher_code'): where.append('a.teacher_code=?'); params.append(request.form['teacher_code'])
    auto_backup()
    with get_db() as conn:
        ids = conn.execute('SELECT a.id FROM attendance_teachers a WHERE ' + ' AND '.join(where), params).fetchall()
        if ids:
            marks = ','.join('?' for _ in ids)
            conn.execute(f'DELETE FROM attendance_teachers WHERE id IN ({marks})', [row['id'] for row in ids])
        conn.commit()
    flash(f'✅ {len(ids)} رکورد غیبت معلم حذف شد.', 'success')
    return _redirect_with_filters('teacher_attendance_report')


def _date_range():
    raw_from, raw_to = request.args.get('date_from', '').strip(), request.args.get('date_to', '').strip()
    try:
        date_from = normalize_jalali_date(raw_from) if raw_from else ''
        date_to = normalize_jalali_date(raw_to) if raw_to else ''
        if date_from and date_to and date_from > date_to:
            raise ValueError('بازه تاریخ نامعتبر است: تاریخ شروع بعد از تاریخ پایان است.')
        return date_from, date_to, None
    except ValueError as exc:
        return raw_from, raw_to, str(exc)


def attendance_report():
    filters, date_error = _report_filters()
    if date_error:
        flash(f'⚠️ {date_error}', 'danger')

    where_full, params_full = _report_where(filters)
    where_base, params_base = _report_where(filters, include_status=False)
    per_page = filters['per_page']
    print_all = request.args.get('print_all') == '1'

    with get_db() as conn:
        analytics_rows = conn.execute(
            f'''SELECT a.student_code, s.first_name, s.last_name, s.class_name, s.grade, a.date, a.status
                FROM attendance_students a JOIN students s ON a.student_code=s.code
                WHERE {where_base}
                ORDER BY a.date''',
            params_base,
        ).fetchall()
        total_row = conn.execute(
            f'''SELECT COUNT(*) AS count
                FROM attendance_students a JOIN students s ON a.student_code=s.code
                WHERE {where_full}''',
            params_full,
        ).fetchone()
        total_records = total_row['count'] or 0
        total_pages = max(1, (total_records + per_page - 1) // per_page)
        page = min(filters['page'], total_pages)
        sort_column = _SORT_COLUMNS[filters['sort']]
        direction = 'ASC' if filters['dir'] == 'asc' else 'DESC'
        order = f'{sort_column} {direction}, a.date DESC, s.last_name ASC, s.first_name ASC'
        records_query = f'''SELECT a.id, s.id AS student_id, s.first_name, s.last_name, s.code, s.grade, s.class_name, s.sida_class,
                                   a.date, a.status, t.first_name || ' ' || t.last_name AS teacher_name
                            FROM attendance_students a JOIN students s ON a.student_code=s.code
                            LEFT JOIN teachers t ON s.teacher_code=t.code
                            WHERE {where_full}
                            ORDER BY {order}'''
        if print_all:
            records = conn.execute(records_query, params_full).fetchall()
        else:
            records = conn.execute(records_query + ' LIMIT ? OFFSET ?',
                                   params_full + [per_page, (page - 1) * per_page]).fetchall()
        grades = conn.execute("SELECT DISTINCT grade FROM students WHERE grade IS NOT NULL AND grade!='' AND (status='فعال' OR status IS NULL) ORDER BY grade").fetchall()
        classes = conn.execute("SELECT DISTINCT class_name FROM students WHERE class_name IS NOT NULL AND class_name!='' AND (status='فعال' OR status IS NULL) ORDER BY class_name").fetchall()
        sida_classes = conn.execute("SELECT DISTINCT sida_class FROM students WHERE sida_class IS NOT NULL AND sida_class!='' AND (status='فعال' OR status IS NULL) ORDER BY sida_class").fetchall()
        teachers_list = conn.execute('SELECT code,first_name,last_name FROM teachers ORDER BY first_name,last_name').fetchall()
        event_dates: dict[str, list[str]] = {}
        if filters['date_from'] and filters['date_to']:
            for event in conn.execute('SELECT date,title FROM events WHERE date>=? AND date<=? ORDER BY date',
                                      (filters['date_from'], filters['date_to'])).fetchall():
                event_dates.setdefault(event['date'], []).append(event['title'])
        removable_absences = None
        if session.get('role') in ('admin', 'manager'):
            removable = conn.execute(
                f'''SELECT COUNT(*) AS count
                    FROM attendance_students a JOIN students s ON a.student_code=s.code
                    WHERE {where_base} AND a.status='غایب' ''',
                params_base,
            ).fetchone()
            removable_absences = removable['count'] or 0
        attendance_layout = load_report_layout(conn, 'attendance_students', session.get('user_id'))
        attendance_layout['visible_count'] = sum(1 for value in attendance_layout['columns'].values() if value)

    if print_all:
        page = 1
        total_pages = 1

    decorated = []
    for record in records:
        item = dict(record)
        item['weekday'] = _jalali_weekday(item['date'])
        item['events'] = list(event_dates.get(item['date'], []))
        item['holiday'] = bool(item['events']) or item['weekday'] == 'جمعه'
        decorated.append(item)

    range_links = [
        {
            'key': key,
            'label': label,
            'active': filters['range'] == key,
            'url': url_for('attendance_report', **_report_query(range=key, date_from=None, date_to=None, page=None)),
        }
        for key, label in _RANGE_LABELS.items() if key != 'custom'
    ]
    sort_links = {}
    for key in _SORT_COLUMNS:
        same = filters['sort'] == key
        if same:
            next_dir = 'desc' if filters['dir'] == 'asc' else 'asc'
        else:
            next_dir = 'desc' if key == 'date' else 'asc'
        sort_links[key] = {
            'url': url_for('attendance_report', **_report_query(sort=key, dir=next_dir, page=None)),
            'state': ('ascending' if filters['dir'] == 'asc' else 'descending') if same else 'none',
        }

    page_numbers = []
    if total_pages > 1:
        window_start = max(1, page - 3)
        window_end = min(total_pages, page + 3)
        if window_start > 1:
            page_numbers.append({'number': 1, 'url': url_for('attendance_report', **_report_query(page=1)), 'active': page == 1})
            if window_start > 2:
                page_numbers.append({'number': None, 'url': None, 'active': False})
        for number in range(window_start, window_end + 1):
            page_numbers.append({'number': number, 'url': url_for('attendance_report', **_report_query(page=number)), 'active': number == page})
        if window_end < total_pages:
            if window_end < total_pages - 1:
                page_numbers.append({'number': None, 'url': None, 'active': False})
            page_numbers.append({'number': total_pages, 'url': url_for('attendance_report', **_report_query(page=total_pages)), 'active': page == total_pages})

    return render_template(
        'attendance_report.html',
        month_groups=_group_records_by_month(decorated),
        analytics=_report_analytics(analytics_rows),
        total_records=total_records,
        page=page,
        total_pages=total_pages,
        per_page=per_page,
        per_page_options=_PER_PAGE_OPTIONS,
        page_numbers=page_numbers,
        prev_url=url_for('attendance_report', **_report_query(page=page - 1)) if page > 1 else None,
        next_url=url_for('attendance_report', **_report_query(page=page + 1)) if page < total_pages else None,
        page_start=(page - 1) * per_page + 1 if total_records else 0,
        page_end=min(total_records, page * per_page),
        grades=grades,
        classes=classes,
        sida_classes=sida_classes,
        teachers_list=teachers_list,
        filters=filters,
        statuses=_STATUSES,
        range_links=range_links,
        sort_links=sort_links,
        filters_description=_filters_description(filters, teachers_list),
        removable_absences=removable_absences,
        print_orientation=attendance_layout['orientation'],
        auto_print=request.args.get('autoprint') == '1',
        query=_report_query(),
        export_url=url_for('attendance_report_export', **_report_query()),
        print_landscape_url=url_for('attendance_report', **_report_query(print='landscape', autoprint='1', print_all='1')),
        print_portrait_url=url_for('attendance_report', **_report_query(print='portrait', autoprint='1', print_all='1')),
        print_layout_url=url_for('report_print_layouts', document='attendance_students'),
        attendance_layout=attendance_layout,
        attendance_print_all=print_all,
    )


def attendance_report_export():
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    filters, _ = _report_filters()
    where_full, params_full = _report_where(filters)
    where_base, params_base = _report_where(filters, include_status=False)

    with get_db() as conn:
        analytics_rows = conn.execute(
            f'''SELECT a.student_code, s.first_name, s.last_name, s.class_name, s.grade, a.date, a.status
                FROM attendance_students a JOIN students s ON a.student_code=s.code
                WHERE {where_base}
                ORDER BY a.date''',
            params_base,
        ).fetchall()
        rows = conn.execute(
            f'''SELECT s.first_name, s.last_name, s.code, s.grade, s.class_name, s.sida_class,
                       a.date, a.status, t.first_name || ' ' || t.last_name AS teacher_name
                FROM attendance_students a JOIN students s ON a.student_code=s.code
                LEFT JOIN teachers t ON s.teacher_code=t.code
                WHERE {where_full}
                ORDER BY a.date DESC, s.last_name, s.first_name''',
            params_full,
        ).fetchall()
        teachers_list = conn.execute('SELECT code,first_name,last_name FROM teachers ORDER BY first_name,last_name').fetchall()

    analytics = _report_analytics(analytics_rows)
    description = _filters_description(filters, teachers_list)
    navy = '145DA0'
    light = 'EAF3FB'
    border_color = 'C8D8E8'
    thin = Side(style='thin', color=border_color)
    header_font = Font(name='Vazirmatn', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='Vazirmatn', size=9, color='1F2937')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def style_header(sheet, row_index, columns):
        for column in range(1, columns + 1):
            cell = sheet.cell(row_index, column)
            cell.fill = PatternFill('solid', fgColor=navy)
            cell.font = header_font
            cell.alignment = center

    def style_body(sheet, first_row, columns):
        for row_index in range(first_row, sheet.max_row + 1):
            for column in range(1, columns + 1):
                cell = sheet.cell(row_index, column)
                cell.font = body_font
                cell.alignment = center
                cell.border = Border(bottom=thin)
            if row_index % 2 == 1:
                for column in range(1, columns + 1):
                    sheet.cell(row_index, column).fill = PatternFill('solid', fgColor='F7FAFC')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'خلاصه'
    ws.sheet_view.rightToLeft = True
    columns = 6
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.cell(1, 1, 'گزارش حضور و غیاب دانش‌آموزان — خلاصه')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.cell(2, 1, f"{current_app.config['SCHOOL_NAME']}  |  تاریخ تهیه: {filters['date_from'] or '—'} تا {filters['date_to'] or '—'}")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)
    ws.cell(3, 1, description)
    ws.cell(5, 1, 'وضعیت')
    ws.cell(5, 2, 'تعداد')
    style_header(ws, 5, 2)
    summary_rows = [
        ('حاضر', analytics['totals']['حاضر']),
        ('غایب', analytics['totals']['غایب']),
        ('مرخصی', analytics['totals']['مرخصی']),
        ('تاخیر', analytics['totals']['تاخیر']),
        ('مجموع ثبت‌شده', analytics['total_marks']),
        ('نرخ حضور', f"{analytics['attendance_rate'] if analytics['attendance_rate'] is not None else '—'}%"),
        ('دانش‌آموزان دارای غیبت', analytics['absent_students']),
        ('دانش‌آموزان دارای تاخیر', analytics['late_students']),
    ]
    for offset, (label, value) in enumerate(summary_rows, 6):
        ws.cell(offset, 1, label)
        ws.cell(offset, 2, value)
    style_body(ws, 6, 2)
    ws.cell(17, 1, 'ماه')
    ws.cell(17, 2, 'مجموع')
    ws.cell(17, 3, 'حاضر')
    ws.cell(17, 4, 'غایب')
    ws.cell(17, 5, 'مرخصی')
    ws.cell(17, 6, 'تاخیر')
    style_header(ws, 17, 6)
    for offset, month in enumerate(analytics['month_trend'], 18):
        ws.cell(offset, 1, month['label'])
        ws.cell(offset, 2, month['total'])
        ws.cell(offset, 3, month['present'])
        ws.cell(offset, 4, month['absent'])
        ws.cell(offset, 5, month['excused'])
        ws.cell(offset, 6, month['late'])
    style_body(ws, 18, 6)
    for column, width in enumerate((22, 14, 12, 12, 12, 12), 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws_students = wb.create_sheet('خلاصه دانش‌آموزان')
    ws_students.sheet_view.rightToLeft = True
    student_headers = ['ردیف', 'نام و نام خانوادگی', 'کد', 'پایه', 'کلاس', 'حاضر', 'غایب', 'تاخیر', 'مرخصی', 'مجموع', 'درصد حضور', 'بیشترین غیبت متوالی']
    for column, header in enumerate(student_headers, 1):
        ws_students.cell(1, column, header)
    style_header(ws_students, 1, len(student_headers))
    for offset, student in enumerate(analytics['students'], 2):
        values = [offset - 1, student['name'], student['code'], student['grade'], student['class_name'],
                  student['present'], student['absent'], student['late'], student['excused'], student['total'],
                  f"{student['rate'] if student['rate'] is not None else '—'}%", student['streak']]
        for column, value in enumerate(values, 1):
            ws_students.cell(offset, column, _excel_value(value))
    style_body(ws_students, 2, len(student_headers))
    for column, width in enumerate((7, 26, 14, 12, 16, 9, 9, 9, 10, 10, 12, 18), 1):
        ws_students.column_dimensions[get_column_letter(column)].width = width
    ws_students.freeze_panes = 'A2'
    ws_students.auto_filter.ref = f'A1:{get_column_letter(len(student_headers))}{max(1, ws_students.max_row)}'

    ws_details = wb.create_sheet('جزئیات')
    ws_details.sheet_view.rightToLeft = True
    detail_headers = ['ردیف', 'نام', 'نام خانوادگی', 'کد', 'پایه', 'کلاس', 'کلاس سیدا', 'معلم', 'تاریخ', 'روز هفته', 'وضعیت', 'مناسبت']
    for column, header in enumerate(detail_headers, 1):
        ws_details.cell(1, column, header)
    style_header(ws_details, 1, len(detail_headers))
    for offset, record in enumerate(rows, 2):
        values = [offset - 1, record['first_name'], record['last_name'], record['code'], record['grade'] or '',
                  record['class_name'] or '', record['sida_class'] or '', record['teacher_name'] or '',
                  record['date'], _jalali_weekday(record['date']), record['status'], '']
        for column, value in enumerate(values, 1):
            ws_details.cell(offset, column, _excel_value(value))
    style_body(ws_details, 2, len(detail_headers))
    for column, width in enumerate((7, 16, 18, 14, 12, 16, 18, 20, 12, 12, 10, 18), 1):
        ws_details.column_dimensions[get_column_letter(column)].width = width
    ws_details.freeze_panes = 'A2'
    ws_details.auto_filter.ref = f'A1:{get_column_letter(len(detail_headers))}{max(1, ws_details.max_row)}'
    ws_details.page_setup.orientation = 'landscape'
    ws_details.page_setup.paperSize = ws_details.PAPERSIZE_A4
    ws_details.sheet_properties.pageSetUpPr.fitToPage = True
    ws_details.page_setup.fitToWidth = 1
    ws_details.page_setup.fitToHeight = 0
    ws_details.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    ws_details.print_title_rows = '1:1'

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='گزارش_حضور_و_غیاب_دانش‌آموزان.xlsx')


def teacher_attendance_report():
    date_from, date_to, date_error = _date_range()
    teacher_code, status_filter = request.args.get('teacher_code', '').strip(), request.args.get('status', '').strip()
    with get_db() as conn:
        query = '''SELECT a.id,t.first_name,t.last_name,t.code,t.subject,t.class_name,a.date,a.status,a.late_duration
                   FROM attendance_teachers a JOIN teachers t ON a.teacher_code=t.code WHERE 1=1'''
        params = []
        if date_from: query += ' AND a.date>=?'; params.append(date_from)
        if date_to: query += ' AND a.date<=?'; params.append(date_to)
        if teacher_code: query += ' AND a.teacher_code=?'; params.append(teacher_code)
        if status_filter: query += ' AND a.status=?'; params.append(status_filter)
        query += ' ORDER BY a.date DESC,t.first_name,t.last_name'
        records = conn.execute(query, params).fetchall()
        summary_query = '''SELECT t.first_name,t.last_name,t.code,
                          SUM(CASE WHEN a.status='غایب' THEN 1 ELSE 0 END) AS total_absent,
                          SUM(CASE WHEN a.status='تاخیر' THEN 1 ELSE 0 END) AS total_late,
                          SUM(CASE WHEN a.status='تاخیر' THEN COALESCE(a.late_duration,0) ELSE 0 END) AS total_late_minutes
                          FROM attendance_teachers a JOIN teachers t ON a.teacher_code=t.code WHERE 1=1'''
        summary_params = []
        if date_from: summary_query += ' AND a.date>=?'; summary_params.append(date_from)
        if date_to: summary_query += ' AND a.date<=?'; summary_params.append(date_to)
        if teacher_code: summary_query += ' AND a.teacher_code=?'; summary_params.append(teacher_code)
        if status_filter: summary_query += ' AND a.status=?'; summary_params.append(status_filter)
        summary_query += ' GROUP BY t.code ORDER BY total_absent DESC,total_late DESC'
        summary = conn.execute(summary_query, summary_params).fetchall()
        teachers_list = conn.execute('SELECT code,first_name,last_name FROM teachers ORDER BY first_name,last_name').fetchall()
        teacher_layout = load_report_layout(conn, 'attendance_teachers', session.get('user_id'))
    if date_error:
        flash(f'⚠️ {date_error}', 'danger')
    return render_template(
        'teacher_attendance_report.html', records=records, summary=summary, teachers_list=teachers_list,
        date_from=date_from, date_to=date_to, teacher_code=teacher_code, status_filter=status_filter,
        teacher_layout=teacher_layout,
        print_orientation=teacher_layout['orientation'],
        auto_print=request.args.get('autoprint') == '1',
        print_landscape_url=url_for('teacher_attendance_report', date_from=date_from or None, date_to=date_to or None,
                                    teacher_code=teacher_code or None, status=status_filter or None,
                                    print='landscape', autoprint='1'),
        print_portrait_url=url_for('teacher_attendance_report', date_from=date_from or None, date_to=date_to or None,
                                   teacher_code=teacher_code or None, status=status_filter or None,
                                   print='portrait', autoprint='1'),
        print_layout_url=url_for('report_print_layouts', document='attendance_teachers'),
    )


def register(app):
    app.add_url_rule('/service_drivers', 'service_drivers', service_drivers, methods=['GET', 'POST'])
    app.add_url_rule('/service_drivers/print', 'service_drivers_print', service_drivers_print)
    app.add_url_rule('/class_management', 'class_management', class_management, methods=['GET', 'POST'])
    app.add_url_rule('/attendance_students', 'attendance_students', attendance_students, methods=['GET', 'POST'])
    app.add_url_rule('/attendance_quick', 'attendance_quick', _attendance_quick, methods=['GET', 'POST'])
    app.add_url_rule('/attendance_teachers', 'attendance_teachers', attendance_teachers, methods=['GET', 'POST'])
    app.add_url_rule('/reports/attendance_report', 'attendance_report', attendance_report)
    app.add_url_rule('/reports/attendance_report/export', 'attendance_report_export', attendance_report_export)
    app.add_url_rule('/reports/teacher_attendance', 'teacher_attendance_report', teacher_attendance_report)
    app.add_url_rule('/attendance_students/<int:id>/delete', 'delete_student_attendance', delete_student_attendance, methods=['POST'])
    app.add_url_rule('/attendance_students/<int:id>/update', 'update_student_attendance', update_student_attendance, methods=['POST'])
    app.add_url_rule('/attendance_teachers/<int:id>/delete', 'delete_teacher_attendance', delete_teacher_attendance, methods=['POST'])
    app.add_url_rule('/reports/attendance_report/delete-absences', 'delete_filtered_student_absences', delete_filtered_student_absences, methods=['POST'])
    app.add_url_rule('/reports/teacher_attendance/delete-absences', 'delete_filtered_teacher_absences', delete_filtered_teacher_absences, methods=['POST'])
