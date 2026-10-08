"""Custom student columns: choose the columns, view them, print A4, export Excel.

The page is independent from the student dashboard, but reuses the same filter
helper and the same teacher scoping, so a teacher only ever sees the students
assigned to their own personnel number and only the columns they may already
read in the student profile.
"""

from __future__ import annotations

from io import BytesIO

from flask import flash, render_template, request, send_file, session

from ..custom_fields import list_fields
from ..database import get_db
from ..security import audit, teacher_scope
from ..student_columns import (EMPTY_LABEL, allowed_keys, build_rows, column_catalogue,
                               default_selection, excel_cell, load_saved_selection,
                               normalise_selection, save_selection, selected_columns)
from .students import (EMPTY, PROFILE_BOOLEAN_FIELDS, PROFILE_LABELS, PROFILE_SECTIONS,
                       filtered_students_query)

FILTER_KEYS = ('q', 'grade', 'class_name', 'sida_class', 'gender', 'teacher_code')
ORIENTATIONS = {'landscape', 'portrait'}
FILE_NAME = 'جدول_ستون_های_دلخواه_دانش_آموزان.xlsx'
WIDE_SELECTION_HINT = 15


def _text(value: object, limit: int = 80) -> str:
    return str(value or '').strip()[:limit]


def _source():
    """POST keeps the whole form; GET keeps the filter query string."""
    return request.form if request.method == 'POST' else request.args


def _filter_values(source) -> dict[str, str]:
    return {key: _text(source.get(key)) for key in FILTER_KEYS}


def _orientation(source) -> str:
    chosen = _text(source.get('orientation'), 20)
    return chosen if chosen in ORIENTATIONS else 'landscape'


def _custom_values(conn, students, columns) -> dict[tuple[int, int], str]:
    """Read the values of the selected custom fields for the selected students."""
    student_ids = [int(row['id']) for row in students]
    field_ids = [int(column['field_id']) for column in columns if column['kind'] == 'custom']
    if not student_ids or not field_ids:
        return {}
    student_marks = ','.join('?' for _ in student_ids)
    field_marks = ','.join('?' for _ in field_ids)
    rows = conn.execute(
        f'''SELECT student_id,field_id,value FROM student_custom_values
            WHERE student_id IN ({student_marks}) AND field_id IN ({field_marks})''',
        student_ids + field_ids,
    ).fetchall()
    return {(int(row['student_id']), int(row['field_id'])): row['value'] or '' for row in rows}


def _filter_options(conn, scope: str | None) -> dict[str, list[dict[str, str]]]:
    """Filter choices, built from the records the current user may see."""
    clauses = ["(s.status = 'فعال' OR s.status IS NULL)"]
    params: list[object] = []
    if scope:
        clauses.append('s.teacher_code = ?')
        params.append(scope)
    where = ' AND '.join(clauses)

    def options(column: str) -> list[dict[str, str]]:
        rows = conn.execute(
            f'''SELECT DISTINCT TRIM(COALESCE({column}, '')) AS value
                FROM students s WHERE {where}
                ORDER BY CASE WHEN value = '' THEN 1 ELSE 0 END, value COLLATE NOCASE''',
            params,
        ).fetchall()
        return [{'value': row['value'] or '__empty__', 'label': row['value'] or EMPTY}
                for row in rows]

    result = {
        'grades': options('s.grade'),
        'classes': options('s.class_name'),
        'sida_classes': options('s.sida_class'),
        'genders': options('s.gender'),
    }
    if scope:
        # A teacher is always limited to their own students, so no teacher list.
        result['teachers'] = []
    else:
        rows = conn.execute(
            """SELECT t.code, TRIM(COALESCE(t.first_name,'') || ' ' || COALESCE(t.last_name,'')) AS name
               FROM teachers t ORDER BY name COLLATE NOCASE"""
        ).fetchall()
        result['teachers'] = [{'value': row['code'], 'label': row['name'] or row['code']}
                              for row in rows if row['code']]
    return result


def _view_context(conn, *, restricted: bool, user_id, persist: bool) -> dict[str, object]:
    """Everything the page, the print sheet and the Excel sheet need."""
    source = _source()
    custom_fields = list_fields(conn, active_only=True)
    groups = column_catalogue(
        conn, labels=PROFILE_LABELS, sections=PROFILE_SECTIONS,
        boolean_fields=PROFILE_BOOLEAN_FIELDS, restricted=restricted, custom_fields=custom_fields,
    )
    allowed = allowed_keys(groups)
    defaults = default_selection(groups)
    previous = load_saved_selection(conn, user_id, allowed, defaults)
    # A GET request shows the saved selection; a POST request is the user's new
    # choice. Only a real POST overwrites the stored preference.
    selected = normalise_selection(source.getlist('columns'), allowed, defaults) if persist else previous
    if persist:
        save_selection(conn, user_id, selected)

    columns = selected_columns(groups, selected)
    scope = teacher_scope()
    query, params, _ = filtered_students_query(source)
    students = conn.execute(query, params).fetchall()
    total_row = conn.execute(
        f"SELECT COUNT(*) AS count FROM students s WHERE (s.status = 'فعال' OR s.status IS NULL)"
        + (' AND s.teacher_code = ?' if scope else ''),
        [scope] if scope else [],
    ).fetchone()

    return {
        'groups': groups,
        'selected': selected,
        'default_keys': defaults,
        'columns': columns,
        'rows': build_rows(students, columns, _custom_values(conn, students, columns), EMPTY_LABEL),
        'student_count': len(students),
        'scope_total': total_row['count'] or 0,
        'filters': _filter_values(source),
        'filter_options': _filter_options(conn, scope),
        'selected_filters': _filter_values(source),
        'changed': persist and selected != previous,
        'saved_count': len(previous),
        'empty_label': EMPTY_LABEL,
    }


def student_columns():
    """Independent page: choose the columns, filter the students, see the table."""
    restricted = session.get('role') == 'teacher'
    user_id = session.get('user_id')
    with get_db() as conn:
        context = _view_context(conn, restricted=restricted, user_id=user_id,
                                persist=request.method == 'POST')
    if context['changed']:
        # Written after the connection committed, so the audit row never waits
        # behind the page's own write lock.
        audit(user_id, 'update_student_columns', 'app_settings', f'user:{user_id}',
              ','.join(context['selected'])[:400])
        flash('ستون‌های دلخواه شما ذخیره شد؛ از این پس همان‌ها نمایش داده میشوند.', 'success')
    return render_template('student_columns.html', **context, restricted=restricted,
                           orientation=_orientation(_source()), wide_hint=WIDE_SELECTION_HINT)


def student_columns_print():
    """A4 sheet of exactly the selected columns for the filtered students."""
    restricted = session.get('role') == 'teacher'
    with get_db() as conn:
        context = _view_context(conn, restricted=restricted, user_id=session.get('user_id'), persist=True)
    return render_template('print_student_columns.html', **context, restricted=restricted,
                           orientation=_orientation(request.form))


def student_columns_export():
    """Excel workbook of the selected columns (numbers stay numeric, A4 ready)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    restricted = session.get('role') == 'teacher'
    orientation = _orientation(request.form)
    with get_db() as conn:
        context = _view_context(conn, restricted=restricted, user_id=session.get('user_id'), persist=True)

    columns = context['columns']
    rows = context['rows']
    headers = ['ردیف'] + [column['label'] for column in columns]
    navy = '145DA0'
    thin = Side(style='thin', color='C8D8E8')
    header_font = Font(name='Vazirmatn', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='Vazirmatn', size=9, color='1F2937')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'ستون‌های دلخواه'
    sheet.sheet_view.rightToLeft = True
    for index, header in enumerate(headers, 1):
        sheet.cell(1, index, header)
    for index in range(1, len(headers) + 1):
        cell = sheet.cell(1, index)
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = header_font
        cell.alignment = center

    for offset, row in enumerate(rows, 2):
        values = [row['index']] + [
            excel_cell(value, numeric=bool(column['numeric']))
            for column, value in zip(columns, row['cells'])
        ]
        for index, value in enumerate(values, 1):
            sheet.cell(offset, index, value)

    # Column widths follow the widest visible text, clamped to a printable range.
    widths = [6]
    for position, column in enumerate(columns):
        longest = max([len(str(column['label']))]
                      + [len(str(row['cells'][position])) for row in rows[:500]])
        widths.append(min(32, max(9, longest + 2)))
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_index, index)
            cell.font = body_font
            cell.alignment = center
            cell.border = Border(bottom=thin)
            if row_index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F7FAFC')

    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}'
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    sheet.print_title_rows = '1:1'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=FILE_NAME)


def register(app):
    app.add_url_rule('/student-columns', endpoint='student_columns', view_func=student_columns,
                     methods=['GET', 'POST'])
    app.add_url_rule('/student-columns/print', endpoint='student_columns_print',
                     view_func=student_columns_print, methods=['POST'])
    app.add_url_rule('/student-columns/export', endpoint='student_columns_export',
                     view_func=student_columns_export, methods=['POST'])