from __future__ import annotations

from io import BytesIO

from flask import current_app, flash, redirect, render_template, request, send_file, session, url_for

from ..database import auto_backup, get_db
from ..dates import get_today_jalali, normalize_digits, normalize_jalali_date, today_string
from ..print_layouts import build_slip_batch, load_layout
from ..security import audit

MONTHS = ['مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند', 'فروردین', 'اردیبهشت', 'خرداد']
SERVICE_TYPES = {'رفت و برگشت', 'رفت', 'برگشت'}
MONTH_ORDER = {month: index for index, month in enumerate(MONTHS, 1)}
SERVICE_STATUS_FILTERS = {
    'debtor': 'بدهکار (پرداخت‌نشده و جزئی)',
    'unpaid': 'پرداخت‌نشده',
    'partial': 'پرداخت جزئی',
    'paid': 'تسویه‌شده',
    'noservice': 'فاقد سرویس',
}
PER_PAGE_OPTIONS = (25, 50, 100, 200)
SORT_OPTIONS = {
    'recent': 'جدیدترین',
    'year': 'سال و ماه',
    'name': 'نام دانش‌آموز',
    'remaining': 'بیشترین مانده',
    'paid': 'بیشترین پرداختی',
}
SORT_COLUMNS = {
    'recent': 'ms.id DESC',
    'year': 'ms.year DESC, ms.id DESC',
    'name': 's.last_name ASC, s.first_name ASC',
    'remaining': '(COALESCE(ms.amount,0) - COALESCE(ms.paid_amount,0)) DESC, ms.id DESC',
    'paid': 'COALESCE(ms.paid_amount,0) DESC, ms.id DESC',
}
SETTINGS_KEYS = ('school_card', 'school_account', 'school_account_name')


def _text(value: object, limit: int = 500) -> str:
    return str(value or '').strip()[:limit]


def _money(value: str | None, label: str) -> tuple[int | None, str | None]:
    raw = normalize_digits(str(value or '')).replace(',', '').replace('٬', '').replace('،', '').strip()
    try:
        number = int(raw or 0)
    except ValueError:
        return None, f'{label} باید عدد صحیح باشد.'
    if number < 0:
        return None, f'{label} نمی‌تواند منفی باشد.'
    return number, None


def _year(value: object, required: bool = False) -> str:
    raw = normalize_digits(_text(value, 10)).strip()
    if not raw:
        if required:
            raise ValueError('سال تحصیلی الزامی است.')
        return ''
    if not raw.isdigit() or not 1390 <= int(raw) <= 1420:
        raise ValueError('سال تحصیلی باید عددی بین ۱۳۹۰ تا ۱۴۲۰ باشد.')
    return raw


def _settings(conn) -> dict[str, str]:
    rows = conn.execute('SELECT key,value FROM app_settings').fetchall()
    stored = {row['key']: row['value'] for row in rows}
    return {key: stored.get(key, '') for key in SETTINGS_KEYS}


def _save_settings(conn, values: dict[str, str]) -> None:
    for key, value in values.items():
        conn.execute('INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',
                     (key, value))



def _date(value: str | None, required: bool = False) -> str:
    raw = _text(value, 20)
    if not raw:
        if required:
            raise ValueError('تاریخ شمسی الزامی است.')
        return ''
    return normalize_jalali_date(raw, required=True)


def _payment_status(amount: int, paid: int) -> str:
    if amount <= 0:
        return 'فاقد سرویس'
    if paid <= 0:
        return 'پرداخت نشده'
    if paid < amount:
        return 'جزئی'
    return 'کامل'


def update_student_service_fees():
    auto_backup()
    with get_db() as conn:
        for key, value in request.form.items():
            if not key.startswith('fee_'):
                continue
            fee, error = _money(value, 'مبلغ سرویس')
            if error:
                flash(error, 'danger')
                return redirect(url_for('monthly_service'))
            conn.execute('UPDATE students SET service_fee=? WHERE code=?', (fee, key.removeprefix('fee_')))
        conn.commit()
    flash('مبالغ اختصاصی سرویس ذخیره شد.', 'success')
    return redirect(url_for('monthly_service'))


def add_student_to_service():
    student_code = _text(request.form.get('student_code'), 50)
    fee, error = _money(request.form.get('fee'), 'مبلغ سرویس')
    if error or not student_code or fee is None or fee <= 0:
        flash(error or 'دانش‌آموز و مبلغ بزرگ‌تر از صفر را انتخاب کنید.', 'danger')
        return redirect(url_for('monthly_service'))
    auto_backup()
    with get_db() as conn:
        if not conn.execute('SELECT code FROM students WHERE code=?', (student_code,)).fetchone():
            flash('دانش‌آموز پیدا نشد.', 'danger')
            return redirect(url_for('monthly_service'))
        conn.execute('UPDATE students SET service_fee=? WHERE code=?', (fee, student_code))
        conn.commit()
    audit(session.get('user_id'), 'add_student_to_service', 'student', student_code)
    flash('دانش‌آموز به فهرست سرویس اضافه شد.', 'success')
    return redirect(url_for('monthly_service'))


def remove_student_service_fee(student_code):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('UPDATE students SET service_fee=0 WHERE code=?', (_text(student_code, 50),))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'remove_student_service_fee', 'student', student_code)
    flash('دانش‌آموز از فهرست سرویس خارج شد.', 'success')
    return redirect(url_for('monthly_service'))


def generate_monthly_services():
    month = _text(request.form.get('month'), 20)
    service_type = _text(request.form.get('service_type', 'رفت و برگشت'), 80)
    try:
        year = _year(request.form.get('year'), required=True)
    except ValueError as exc:
        flash(f'⚠️ {exc}', 'danger')
        return redirect(url_for('monthly_service'))
    if month not in MONTHS or service_type not in SERVICE_TYPES:
        flash('ماه یا نوع سرویس معتبر نیست.', 'danger')
        return redirect(url_for('monthly_service'))
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute(
            '''INSERT OR IGNORE INTO monthly_service(student_code,year,month,service_type,amount,paid_amount,status)
               SELECT s.code,?,?,?,s.service_fee,0,'پرداخت نشده'
               FROM students s
               WHERE (s.status='فعال' OR s.status IS NULL) AND s.service_fee>0
                 AND NOT EXISTS (SELECT 1 FROM monthly_service ms
                                 WHERE ms.student_code=s.code AND ms.year=? AND ms.month=?)''',
            (year, month, service_type, year, month),
        )
        count = cursor.rowcount if cursor.rowcount and cursor.rowcount > 0 else 0
        conn.commit()
    audit(session.get('user_id'), 'generate_monthly_services', 'monthly_service', f'{year}/{month}', str(count))
    flash(f'برای {count} دانش‌آموز سرویس ماهانه صادر شد.', 'success')
    return redirect(url_for('monthly_service'))


def delete_monthly_service(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM monthly_service WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_monthly_service', 'monthly_service', id)
        flash('رکورد سرویس حذف شد.', 'success')
    else:
        flash('رکورد سرویس پیدا نشد.', 'warning')
    return redirect(url_for('monthly_service'))


def delete_zero_monthly_services():
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM monthly_service WHERE COALESCE(amount,0)<=0 AND COALESCE(paid_amount,0)<=0')
        count = cursor.rowcount
        conn.commit()
    flash(f'{count} رکورد بدون مبلغ حذف شد.', 'success')
    return redirect(url_for('monthly_service'))


def financial_help():
    if request.method == 'POST':
        amount, error = _money(request.form.get('amount'), 'مبلغ کمک')
        student_code = _text(request.form.get('student_code'), 50)
        try:
            date = _date(request.form.get('date'))
        except ValueError as exc:
            date, error = '', str(exc)
        if error or not student_code:
            flash(error or 'دانش‌آموز را انتخاب کنید.', 'danger')
            return redirect(url_for('financial_help'))
        auto_backup()
        with get_db() as conn:
            if not conn.execute('SELECT code FROM students WHERE code=?', (student_code,)).fetchone():
                flash('دانش‌آموز پیدا نشد.', 'danger')
                return redirect(url_for('financial_help'))
            conn.execute('INSERT INTO financial_help(student_code,amount,date,description) VALUES(?,?,?,?)',
                         (student_code, amount, date, _text(request.form.get('description'))))
            conn.commit()
        audit(session.get('user_id'), 'create_financial_help', 'financial_help', student_code)
        flash('کمک مالی ثبت شد.', 'success')
        return redirect(url_for('financial_help'))
    with get_db() as conn:
        helps = conn.execute('''SELECT fh.*,s.first_name,s.last_name,s.code FROM financial_help fh
                                JOIN students s ON fh.student_code=s.code ORDER BY fh.date DESC,fh.id DESC''').fetchall()
        students = conn.execute("SELECT first_name,last_name,code FROM students WHERE status='فعال' OR status IS NULL ORDER BY first_name,last_name").fetchall()
    return render_template('financial_help.html', helps=helps, students=students, today=today_string())


def delete_financial_help(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM financial_help WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_financial_help', 'financial_help', id)
        flash('کمک مالی حذف شد.', 'success')
    else:
        flash('رکورد کمک مالی پیدا نشد.', 'warning')
    return redirect(url_for('financial_help'))


def _validate_payment_form(form):
    student_code = _text(form.get('student_code'), 50)
    month = _text(form.get('month'), 20)
    service_type = _text(form.get('service_type', 'رفت و برگشت'), 80)
    amount, amount_error = _money(form.get('amount'), 'مبلغ مصوب')
    paid, paid_error = _money(form.get('paid_amount'), 'مبلغ پرداختی')
    errors = [e for e in (amount_error, paid_error) if e]
    try:
        year = _year(form.get('year'), required=True)
    except ValueError as exc:
        year = _text(form.get('year'), 10)
        errors.append(str(exc))
    if not student_code or month not in MONTHS:
        errors.append('دانش‌آموز و ماه معتبر الزامی است.')
    if service_type not in SERVICE_TYPES:
        errors.append('نوع سرویس معتبر نیست.')
    if amount is None or paid is None:
        errors.append('مبالغ الزامی هستند.')
    if amount is not None and paid is not None and paid > amount:
        errors.append('مبلغ پرداختی نمی‌تواند بیشتر از مبلغ مصوب باشد.')
    try:
        payment_date = _date(form.get('payment_date'))
    except ValueError as exc:
        payment_date = ''
        errors.append(str(exc))
    if paid and paid > 0 and not payment_date:
        errors.append('برای مبلغ پرداختی بزرگ‌تر از صفر، تاریخ پرداخت الزامی است.')
    return {
        'student_code': student_code, 'year': year, 'month': month, 'service_type': service_type,
        'amount': amount if amount is not None else 0, 'paid_amount': paid if paid is not None else 0,
        'payment_date': payment_date, 'description': _text(form.get('description')),
    }, errors


def _service_filters() -> dict[str, object]:
    year = normalize_digits(_text(request.args.get('year'), 10)).strip()
    if year and (not year.isdigit() or not 1390 <= int(year) <= 1420):
        year = ''
    month = _text(request.args.get('month'), 20)
    if month not in MONTHS:
        month = ''
    status = _text(request.args.get('status'), 20)
    if status not in SERVICE_STATUS_FILTERS:
        status = ''
    service_type = _text(request.args.get('service_type'), 80)
    if service_type not in SERVICE_TYPES:
        service_type = ''
    sort = _text(request.args.get('sort'), 20)
    if sort not in SORT_OPTIONS:
        sort = 'recent'
    try:
        per_page = int(request.args.get('per_page', 50))
    except (TypeError, ValueError):
        per_page = 50
    if per_page not in PER_PAGE_OPTIONS:
        per_page = 50
    try:
        page = max(1, int(request.args.get('page', 1)))
    except (TypeError, ValueError):
        page = 1
    return {
        'year': year,
        'month': month,
        'status': status,
        'service_type': service_type,
        'q': _text(request.args.get('q'), 80),
        'class_name': _text(request.args.get('class_name'), 60),
        'sort': sort,
        'per_page': per_page,
        'page': page,
    }


def _service_where(filters: dict[str, object]) -> tuple[str, list[object]]:
    clauses: list[str] = []
    params: list[object] = []
    if filters['year']:
        clauses.append('ms.year=?')
        params.append(filters['year'])
    if filters['month']:
        clauses.append('ms.month=?')
        params.append(filters['month'])
    if filters['service_type']:
        clauses.append('ms.service_type=?')
        params.append(filters['service_type'])
    if filters['class_name']:
        clauses.append('s.class_name=?')
        params.append(filters['class_name'])
    if filters['q']:
        like = f"%{filters['q']}%"
        clauses.append('(s.first_name LIKE ? OR s.last_name LIKE ? OR s.code LIKE ?)')
        params.extend([like, like, like])
    status_clauses = {
        'debtor': "ms.status IN ('پرداخت نشده','جزئی') AND COALESCE(ms.amount,0)>0",
        'unpaid': "ms.status='پرداخت نشده' AND COALESCE(ms.amount,0)>0",
        'partial': "ms.status='جزئی'",
        'paid': "ms.status='کامل'",
        'noservice': "COALESCE(ms.amount,0)<=0",
    }
    if filters['status'] in status_clauses:
        clauses.append(status_clauses[filters['status']])
    return (' AND '.join(clauses) if clauses else '1=1'), params


def _service_query(**overrides) -> dict[str, object]:
    query: dict[str, object] = {}
    for key in ('year', 'month', 'status', 'service_type', 'q', 'class_name', 'sort', 'per_page'):
        value = request.args.get(key)
        if value:
            query[key] = value
    for key, value in overrides.items():
        if value in (None, '', []):
            query.pop(key, None)
        else:
            query[key] = value
    return query


def _service_summary(conn, where: str, params: list[object]) -> dict[str, object]:
    row = conn.execute(
        f'''SELECT COUNT(*) AS records,
                   COALESCE(SUM(ms.amount),0) AS billed,
                   COALESCE(SUM(ms.paid_amount),0) AS paid,
                   SUM(CASE WHEN ms.status IN ('پرداخت نشده','جزئی') AND COALESCE(ms.amount,0)>0 THEN 1 ELSE 0 END) AS debtors,
                   SUM(CASE WHEN COALESCE(ms.amount,0)<=0 THEN 1 ELSE 0 END) AS zero_records
            FROM monthly_service ms JOIN students s ON ms.student_code=s.code
            WHERE {where}''',
        params,
    ).fetchone()
    billed = row['billed'] or 0
    paid = row['paid'] or 0
    debtors = row['debtors'] or 0
    remaining = max(0, billed - paid)
    return {
        'records': row['records'] or 0,
        'billed': billed,
        'paid': paid,
        'remaining': remaining,
        'debtors': debtors,
        'zero_records': row['zero_records'] or 0,
        'collection_rate': round(paid * 100 / billed) if billed else None,
        'average_remaining': round(remaining / debtors) if debtors else 0,
    }


def _service_trend(conn, where: str, params: list[object]) -> list[dict[str, object]]:
    rows = conn.execute(
        f'''SELECT ms.year, ms.month, COUNT(*) AS records,
                   COALESCE(SUM(ms.amount),0) AS billed,
                   COALESCE(SUM(ms.paid_amount),0) AS paid,
                   SUM(CASE WHEN ms.status IN ('پرداخت نشده','جزئی') AND COALESCE(ms.amount,0)>0 THEN 1 ELSE 0 END) AS debtors
            FROM monthly_service ms JOIN students s ON ms.student_code=s.code
            WHERE {where}
            GROUP BY ms.year, ms.month''',
        params,
    ).fetchall()
    items = sorted((dict(row) for row in rows), key=lambda item: (item['year'] or '', MONTH_ORDER.get(item['month'], 99)))
    for item in items:
        item['label'] = f"{item['month']} {item['year']}"
        item['remaining'] = max(0, item['billed'] - item['paid'])
        item['rate'] = round(item['paid'] * 100 / item['billed']) if item['billed'] else None
    return items


def _service_classes(conn, where: str, params: list[object]) -> list[dict[str, object]]:
    rows = conn.execute(
        f'''SELECT COALESCE(NULLIF(TRIM(s.class_name),''),'ثبت نشده') AS class_name,
                   COUNT(*) AS records,
                   COALESCE(SUM(ms.amount),0) AS billed,
                   COALESCE(SUM(ms.paid_amount),0) AS paid,
                   SUM(CASE WHEN ms.status IN ('پرداخت نشده','جزئی') AND COALESCE(ms.amount,0)>0 THEN 1 ELSE 0 END) AS debtors
            FROM monthly_service ms JOIN students s ON ms.student_code=s.code
            WHERE {where}
            GROUP BY class_name
            ORDER BY (COALESCE(SUM(ms.amount),0) - COALESCE(SUM(ms.paid_amount),0)) DESC''',
        params,
    ).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        item['remaining'] = max(0, item['billed'] - item['paid'])
        item['rate'] = round(item['paid'] * 100 / item['billed']) if item['billed'] else None
        items.append(item)
    return items[:15]


def _service_top_debtors(conn, where: str, params: list[object]) -> list[dict[str, object]]:
    rows = conn.execute(
        f'''SELECT s.first_name, s.last_name, s.code,
                   COALESCE(NULLIF(TRIM(s.class_name),''),'ثبت نشده') AS class_name,
                   COUNT(*) AS records,
                   COALESCE(SUM(ms.amount),0) AS billed,
                   COALESCE(SUM(ms.paid_amount),0) AS paid
            FROM monthly_service ms JOIN students s ON ms.student_code=s.code
            WHERE {where} AND ms.status IN ('پرداخت نشده','جزئی') AND COALESCE(ms.amount,0)>0
            GROUP BY s.code
            ORDER BY (COALESCE(SUM(ms.amount),0) - COALESCE(SUM(ms.paid_amount),0)) DESC
            LIMIT 8''',
        params,
    ).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        item['remaining'] = max(0, item['billed'] - item['paid'])
        items.append(item)
    return items


def monthly_service():
    if request.method == 'POST':
        data, errors = _validate_payment_form(request.form)
        if errors:
            for error in errors:
                flash(error, 'danger')
            return redirect(url_for('monthly_service'))
        auto_backup()
        with get_db() as conn:
            if not conn.execute('SELECT code FROM students WHERE code=?', (data['student_code'],)).fetchone():
                flash('دانش‌آموز پیدا نشد.', 'danger')
                return redirect(url_for('monthly_service'))
            existing = conn.execute(
                'SELECT id FROM monthly_service WHERE student_code=? AND year=? AND month=?',
                (data['student_code'], data['year'], data['month']),
            ).fetchone()
            if existing:
                flash('⚠️ برای این دانش‌آموز در این ماه رکورد وجود دارد؛ همان رکورد را ویرایش کنید.', 'warning')
                return redirect(url_for('edit_payment', id=existing['id']))
            conn.execute('''INSERT INTO monthly_service(student_code,year,month,service_type,amount,paid_amount,payment_date,status,description)
                            VALUES(?,?,?,?,?,?,?,?,?)''',
                         (data['student_code'], data['year'], data['month'], data['service_type'], data['amount'], data['paid_amount'],
                          data['payment_date'], _payment_status(data['amount'], data['paid_amount']), data['description']))
            conn.commit()
        audit(session.get('user_id'), 'create_payment', 'monthly_service', data['student_code'])
        flash('رکورد پرداخت با موفقیت ثبت شد.', 'success')
        return redirect(url_for('monthly_service'))

    filters = _service_filters()
    where, params = _service_where(filters)
    per_page = filters['per_page']
    page = filters['page']
    with get_db() as conn:
        total_row = conn.execute(
            f'SELECT COUNT(*) AS count FROM monthly_service ms JOIN students s ON ms.student_code=s.code WHERE {where}',
            params,
        ).fetchone()
        total_records = total_row['count'] or 0
        total_pages = max(1, (total_records + per_page - 1) // per_page)
        page = min(page, total_pages)
        services = conn.execute(
            f'''SELECT ms.*, s.first_name, s.last_name, s.code, s.grade, s.class_name
                FROM monthly_service ms JOIN students s ON ms.student_code=s.code
                WHERE {where}
                ORDER BY {SORT_COLUMNS[filters['sort']]}
                LIMIT ? OFFSET ?''',
            params + [per_page, (page - 1) * per_page],
        ).fetchall()
        students = conn.execute(
            "SELECT first_name,last_name,code,service_fee,class_name FROM students WHERE status='فعال' OR status IS NULL ORDER BY first_name,last_name"
        ).fetchall()
        summary = _service_summary(conn, where, params)
        trend = _service_trend(conn, where, params)
        classes_summary = _service_classes(conn, where, params)
        top_debtors = _service_top_debtors(conn, where, params)
        years = conn.execute("SELECT DISTINCT year FROM monthly_service WHERE TRIM(COALESCE(year,''))<>'' ORDER BY year DESC").fetchall()
        class_options = conn.execute(
            "SELECT DISTINCT class_name FROM students WHERE TRIM(COALESCE(class_name,''))<>'' AND (status='فعال' OR status IS NULL) ORDER BY class_name"
        ).fetchall()
        zero_total_row = conn.execute(
            'SELECT COUNT(*) AS count FROM monthly_service WHERE COALESCE(amount,0)<=0 AND COALESCE(paid_amount,0)<=0'
        ).fetchone()
        settings = _settings(conn)

    page_numbers = []
    if total_pages > 1:
        window_start = max(1, page - 3)
        window_end = min(total_pages, page + 3)
        if window_start > 1:
            page_numbers.append({'number': 1, 'url': url_for('monthly_service', **_service_query(page=1)), 'active': page == 1})
            if window_start > 2:
                page_numbers.append({'number': None, 'url': None, 'active': False})
        for number in range(window_start, window_end + 1):
            page_numbers.append({'number': number, 'url': url_for('monthly_service', **_service_query(page=number)), 'active': number == page})
        if window_end < total_pages:
            if window_end < total_pages - 1:
                page_numbers.append({'number': None, 'url': None, 'active': False})
            page_numbers.append({'number': total_pages, 'url': url_for('monthly_service', **_service_query(page=total_pages)), 'active': page == total_pages})

    return render_template(
        'monthly_service.html',
        services=services,
        students=students,
        today=today_string(),
        filters=filters,
        summary=summary,
        trend=trend,
        classes_summary=classes_summary,
        top_debtors=top_debtors,
        years=[row['year'] for row in years],
        months=MONTHS,
        service_types=sorted(SERVICE_TYPES),
        status_filters=SERVICE_STATUS_FILTERS,
        sort_options=SORT_OPTIONS,
        per_page_options=PER_PAGE_OPTIONS,
        class_options=[row['class_name'] for row in class_options],
        settings=settings,
        zero_total=zero_total_row['count'] or 0,
        total_records=total_records,
        page=page,
        total_pages=total_pages,
        per_page=per_page,
        page_numbers=page_numbers,
        prev_url=url_for('monthly_service', **_service_query(page=page - 1)) if page > 1 else None,
        next_url=url_for('monthly_service', **_service_query(page=page + 1)) if page < total_pages else None,
        page_start=(page - 1) * per_page + 1 if total_records else 0,
        page_end=min(total_records, page * per_page),
        export_url=url_for('monthly_service_export', **_service_query()),
        print_both_url=url_for('monthly_service_print', **_service_query(kind='both')),
        print_requests_url=url_for('monthly_service_print', **_service_query(kind='requests')),
        print_receipts_url=url_for('monthly_service_print', **_service_query(kind='receipts')),
    )


def payment_history(student_code):
    with get_db() as conn:
        student = conn.execute('SELECT first_name,last_name,code,grade,class_name,sida_class FROM students WHERE code=?', (_text(student_code, 50),)).fetchone()
        if not student:
            flash('دانش‌آموز پیدا نشد.', 'danger')
            return redirect(url_for('monthly_service'))
        history = conn.execute('SELECT * FROM monthly_service WHERE student_code=? ORDER BY year DESC,month DESC,id DESC', (student_code,)).fetchall()
    return render_template('payment_history.html', history=history, student=student)


def edit_payment(id):
    with get_db() as conn:
        payment = conn.execute('SELECT ms.*,s.first_name,s.last_name,s.code FROM monthly_service ms JOIN students s ON ms.student_code=s.code WHERE ms.id=?', (id,)).fetchone()
        students = conn.execute("SELECT first_name,last_name,code FROM students WHERE status='فعال' OR status IS NULL ORDER BY first_name,last_name").fetchall()
    if not payment:
        flash('رکورد پرداخت پیدا نشد.', 'danger')
        return redirect(url_for('monthly_service'))
    if request.method == 'POST':
        data, errors = _validate_payment_form(request.form)
        payment_view = dict(payment)
        payment_view.update(data)
        if errors:
            for error in errors:
                flash(error, 'danger')
            return render_template('edit_payment.html', payment=payment_view, students=students)
        auto_backup()
        with get_db() as conn:
            if not conn.execute('SELECT code FROM students WHERE code=?', (data['student_code'],)).fetchone():
                flash('دانش‌آموز انتخاب‌شده معتبر نیست.', 'danger')
                return render_template('edit_payment.html', payment=payment_view, students=students)
            conn.execute('''UPDATE monthly_service SET student_code=?,year=?,month=?,service_type=?,amount=?,paid_amount=?,
                            payment_date=?,status=?,description=? WHERE id=?''',
                         (data['student_code'], data['year'], data['month'], data['service_type'], data['amount'], data['paid_amount'],
                          data['payment_date'], _payment_status(data['amount'], data['paid_amount']), data['description'], id))
            conn.commit()
        audit(session.get('user_id'), 'edit_payment', 'monthly_service', id)
        flash('تمام اطلاعات پرداخت به‌روزرسانی شد.', 'success')
        return redirect(url_for('monthly_service'))
    return render_template('edit_payment.html', payment=payment, students=students)


def debtors_by_month():
    month = _text(request.args.get('month'), 20)
    year = _text(request.args.get('year'), 10)
    query = '''SELECT s.first_name,s.last_name,ms.year,ms.month,ms.amount,ms.paid_amount,
               (ms.amount-ms.paid_amount) AS remaining,ms.service_type
               FROM monthly_service ms JOIN students s ON ms.student_code=s.code
               WHERE ms.status IN ('پرداخت نشده','جزئی') AND ms.amount>0'''
    params: list[object] = []
    if month:
        query += ' AND ms.month=?'; params.append(month)
    if year:
        query += ' AND ms.year=?'; params.append(year)
    query += ' ORDER BY ms.year,ms.month'
    with get_db() as conn:
        debtors = conn.execute(query, params).fetchall()
        years = conn.execute('SELECT DISTINCT year FROM monthly_service ORDER BY year DESC').fetchall()
    return render_template('debtors_report.html', debtors=debtors, months=MONTHS, years=years, month=month, year=year)


def _filters_description(filters: dict[str, object]) -> str:
    parts: list[str] = []
    if filters['year']:
        parts.append(f"سال: {filters['year']}")
    if filters['month']:
        parts.append(f"ماه: {filters['month']}")
    if filters['status']:
        parts.append(f"وضعیت: {SERVICE_STATUS_FILTERS[filters['status']]}")
    if filters['service_type']:
        parts.append(f"نوع: {filters['service_type']}")
    if filters['class_name']:
        parts.append(f"کلاس: {filters['class_name']}")
    if filters['q']:
        parts.append(f"جست‌وجو: {filters['q']}")
    return ' | '.join(parts) or 'همه رکوردها'


def save_monthly_service_settings():
    values = {
        'school_card': _text(request.form.get('school_card'), 40),
        'school_account': _text(request.form.get('school_account'), 40),
        'school_account_name': _text(request.form.get('school_account_name'), 120),
    }
    auto_backup()
    with get_db() as conn:
        _save_settings(conn, values)
        conn.commit()
    audit(session.get('user_id'), 'save_service_settings', 'app_settings', 'monthly_service')
    flash('مشخصات بانکی آموزشگاه ذخیره شد.', 'success')
    return redirect(url_for('monthly_service'))


def rebuild_monthly_status():
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute(
            '''UPDATE monthly_service SET status = CASE
                   WHEN COALESCE(amount,0)<=0 THEN 'فاقد سرویس'
                   WHEN COALESCE(paid_amount,0)<=0 THEN 'پرداخت نشده'
                   WHEN COALESCE(paid_amount,0)<COALESCE(amount,0) THEN 'جزئی'
                   ELSE 'کامل' END''')
        count = cursor.rowcount or 0
        conn.commit()
    audit(session.get('user_id'), 'rebuild_monthly_status', 'monthly_service', '', f'{count} رکورد')
    flash(f'وضعیت {count} رکورد بازسازی شد.', 'success')
    return redirect(url_for('monthly_service'))


def monthly_service_export():
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    filters = _service_filters()
    academic_months = MONTHS
    today = get_today_jalali()
    annual_year = filters['year'] or str(today[0] if today[1] >= 7 else today[0] - 1)
    where, params = _service_where(filters)
    with get_db() as conn:
        rows = conn.execute(
            f'''SELECT ms.id, s.first_name, s.last_name, s.code, s.grade, s.class_name,
                       ms.year, ms.month, ms.service_type, ms.amount, ms.paid_amount,
                       ms.payment_date, ms.status, ms.description
                FROM monthly_service ms JOIN students s ON ms.student_code=s.code
                WHERE {where}
                ORDER BY {SORT_COLUMNS[filters['sort']]}''',
            params,
        ).fetchall()
        summary = _service_summary(conn, where, params)
        trend = _service_trend(conn, where, params)
        classes = _service_classes(conn, where, params)
        annual_month_placeholders = ','.join('?' for _ in academic_months)
        annual_students = conn.execute(
            f'''SELECT s.first_name, s.last_name, s.code, s.grade, s.class_name,
                       COALESCE(s.service_fee,0) AS service_fee
                FROM students s
                WHERE s.status='فعال' OR s.status IS NULL OR TRIM(s.status)=''
                   OR EXISTS (
                       SELECT 1 FROM monthly_service ms
                       WHERE ms.student_code=s.code AND ms.year=?
                         AND ms.month IN ({annual_month_placeholders})
                   )
                ORDER BY s.last_name ASC, s.first_name ASC''',
            [annual_year, *academic_months],
        ).fetchall()
        annual_service_rows = conn.execute(
            f'''SELECT ms.student_code, ms.month, ms.amount, ms.paid_amount
                FROM monthly_service ms
                WHERE ms.year=? AND ms.month IN ({annual_month_placeholders})
                ORDER BY ms.id ASC''',
            [annual_year, *academic_months],
        ).fetchall()

    today_label = f'{today[0]:04d}/{today[1]:02d}/{today[2]:02d}'
    navy = '145DA0'
    border_color = 'C8D8E8'
    thin = Side(style='thin', color=border_color)
    header_font = Font(name='Vazirmatn', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='Vazirmatn', size=9, color='1F2937')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def style_header(sheet, row_index, columns):
        for column in range(1, columns + 1):
            cell = sheet.cell(row_index, column)
            cell.fill = PatternFill('solid', fgColor=navy)
            cell.font = header_font
            cell.alignment = center

    def style_body(sheet, first_row, columns):
        for row_index in range(first_row, sheet.max_row + 1):
            for column in range(1, columns + 1):
                cell = sheet.cell(row_index, column)
                cell.font = body_font
                cell.alignment = center
                cell.border = Border(bottom=thin)
            if row_index % 2 == 1:
                for column in range(1, columns + 1):
                    sheet.cell(row_index, column).fill = PatternFill('solid', fgColor='F7FAFC')

    annual_records = {}
    for record in annual_service_rows:
        annual_records.setdefault(record['student_code'], {})[record['month']] = record

    annual_rows = []
    for student in annual_students:
        monthly_fee = int(student['service_fee'] or 0)
        records_by_month = annual_records.get(student['code'], {})
        statuses = []
        total_amount = 0
        total_paid = 0
        missing_months = 0
        for month in academic_months:
            record = records_by_month.get(month)
            if record:
                amount = int(record['amount'] or 0)
                paid = int(record['paid_amount'] or 0)
                total_amount += amount
                total_paid += paid
                if amount <= 0:
                    status = 'فاقد سرویس'
                elif paid <= 0:
                    status = 'پرداخت نشده'
                elif paid < amount:
                    status = 'پرداخت جزئی'
                else:
                    status = 'تسویه کامل'
            elif monthly_fee > 0:
                status = 'ثبت نشده'
                missing_months += 1
            else:
                status = 'فاقد سرویس'
            statuses.append(status)

        remaining = max(0, total_amount - total_paid)
        if not records_by_month and monthly_fee <= 0:
            overall_status = 'فاقد سرویس'
        elif missing_months:
            overall_status = 'ناقص؛ ماه ثبت‌نشده'
        elif total_amount <= 0:
            overall_status = 'فاقد سرویس'
        elif total_paid <= 0:
            overall_status = 'پرداخت نشده'
        elif total_paid < total_amount:
            overall_status = 'پرداخت جزئی'
        else:
            overall_status = 'تسویه کامل'

        annual_rows.append([
            len(annual_rows) + 1,
            f"{student['first_name'] or ''} {student['last_name'] or ''}".strip(),
            student['code'],
            student['grade'] or '',
            student['class_name'] or '',
            monthly_fee,
            *statuses,
            total_amount,
            total_paid,
            remaining,
            overall_status,
        ])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'خلاصه'
    ws.sheet_view.rightToLeft = True
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.cell(1, 1, 'گزارش سرویس ماهانه — خلاصه')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    ws.cell(2, 1, f"{current_app.config['SCHOOL_NAME']}  |  تاریخ تهیه: {today_label}  |  {_filters_description(filters)}")
    headers = ['شرح', 'مقدار']
    for column, header in enumerate(headers, 1):
        ws.cell(4, column, header)
    style_header(ws, 4, 2)
    summary_rows = [
        ('تعداد رکورد', summary['records']),
        ('مجموع مطالبات (تومان)', summary['billed']),
        ('کل واریزی‌ها (تومان)', summary['paid']),
        ('مانده بدهی (تومان)', summary['remaining']),
        ('تعداد بدهکاران', summary['debtors']),
        ('میانگین مانده هر بدهکار (تومان)', summary['average_remaining']),
        ('نرخ وصول', f"{summary['collection_rate'] if summary['collection_rate'] is not None else '—'}%"),
    ]
    for offset, (label, value) in enumerate(summary_rows, 5):
        ws.cell(offset, 1, label)
        ws.cell(offset, 2, value)
    style_body(ws, 5, 2)
    trend_start = 5 + len(summary_rows) + 2
    trend_headers = ['دوره', 'رکورد', 'مطالبات', 'واریزی', 'مانده', 'نرخ وصول']
    for column, header in enumerate(trend_headers, 1):
        ws.cell(trend_start, column, header)
    style_header(ws, trend_start, len(trend_headers))
    for offset, item in enumerate(trend, trend_start + 1):
        values = [item['label'], item['records'], item['billed'], item['paid'], item['remaining'],
                  f"{item['rate'] if item['rate'] is not None else '—'}%"]
        for column, value in enumerate(values, 1):
            ws.cell(offset, column, value)
    style_body(ws, trend_start + 1, len(trend_headers))
    class_start = trend_start + len(trend) + 3
    class_headers = ['کلاس', 'رکورد', 'مطالبات', 'واریزی', 'مانده', 'بدهکاران', 'نرخ وصول']
    for column, header in enumerate(class_headers, 1):
        ws.cell(class_start, column, header)
    style_header(ws, class_start, len(class_headers))
    for offset, item in enumerate(classes, class_start + 1):
        values = [item['class_name'], item['records'], item['billed'], item['paid'], item['remaining'],
                  item['debtors'], f"{item['rate'] if item['rate'] is not None else '—'}%"]
        for column, value in enumerate(values, 1):
            ws.cell(offset, column, value)
    style_body(ws, class_start + 1, len(class_headers))
    for column, width in enumerate((34, 14, 18, 18, 18, 14), 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws_annual = wb.create_sheet('وضعیت سالانه', 1)
    ws_annual.sheet_view.rightToLeft = True
    annual_headers = [
        'ردیف', 'دانش‌آموز', 'کد', 'پایه', 'کلاس', 'هزینه ماهانه',
        *academic_months, 'جمع مبلغ مصوب', 'جمع پرداختی', 'مانده', 'وضعیت کلی',
    ]
    annual_last_column = len(annual_headers)
    ws_annual.merge_cells(start_row=1, start_column=1, end_row=1, end_column=annual_last_column)
    ws_annual.cell(1, 1, 'گزارش سالانه وضعیت پرداخت سرویس دانش‌آموزان')
    ws_annual.merge_cells(start_row=2, start_column=1, end_row=2, end_column=annual_last_column)
    ws_annual.cell(
        2, 1,
        f"{current_app.config['SCHOOL_NAME']}  |  سال تحصیلی {annual_year}-{int(annual_year) + 1}  |  دوره مهر تا خرداد  |  تاریخ تهیه: {today_label}",
    )
    for column, header in enumerate(annual_headers, 1):
        ws_annual.cell(4, column, header)
    style_header(ws_annual, 4, annual_last_column)
    for row_index, values in enumerate(annual_rows, 5):
        for column, value in enumerate(values, 1):
            ws_annual.cell(row_index, column, _excel_value(value))
    style_body(ws_annual, 5, annual_last_column)
    annual_status_fills = {
        'تسویه کامل': 'E8F5E9',
        'پرداخت جزئی': 'FFF4D6',
        'پرداخت نشده': 'FDECEC',
        'ثبت نشده': 'FDECEC',
        'ناقص؛ ماه ثبت‌نشده': 'FFF4D6',
        'فاقد سرویس': 'F3F4F6',
    }
    for row_index in range(5, ws_annual.max_row + 1):
        for column in range(7, 7 + len(academic_months)):
            cell = ws_annual.cell(row_index, column)
            fill_color = annual_status_fills.get(str(cell.value))
            if fill_color:
                cell.fill = PatternFill('solid', fgColor=fill_color)
        overall_cell = ws_annual.cell(row_index, annual_last_column)
        fill_color = annual_status_fills.get(str(overall_cell.value))
        if fill_color:
            overall_cell.fill = PatternFill('solid', fgColor=fill_color)
    annual_widths = (7, 24, 14, 10, 16, 14) + (14,) * len(academic_months) + (16, 16, 16, 24)
    for column, width in enumerate(annual_widths, 1):
        ws_annual.column_dimensions[get_column_letter(column)].width = width
    ws_annual.freeze_panes = 'A5'
    ws_annual.auto_filter.ref = f'A4:{get_column_letter(annual_last_column)}{max(4, ws_annual.max_row)}'
    ws_annual.page_setup.orientation = 'landscape'
    ws_annual.page_setup.paperSize = ws_annual.PAPERSIZE_A4
    ws_annual.sheet_properties.pageSetUpPr.fitToPage = True
    ws_annual.page_setup.fitToWidth = 1
    ws_annual.page_setup.fitToHeight = 0
    ws_annual.page_margins = PageMargins(left=0.2, right=0.2, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    ws_annual.print_title_rows = '1:4'

    ws_records = wb.create_sheet('رکوردها')
    ws_records.sheet_view.rightToLeft = True
    record_headers = ['ردیف', 'دانش‌آموز', 'کد', 'پایه', 'کلاس', 'سال', 'ماه', 'نوع سرویس',
                      'مبلغ مصوب', 'پرداختی', 'مانده', 'وضعیت', 'تاریخ پرداخت', 'توضیحات']
    for column, header in enumerate(record_headers, 1):
        ws_records.cell(1, column, header)
    style_header(ws_records, 1, len(record_headers))
    for offset, record in enumerate(rows, 2):
        amount = int(record['amount'] or 0)
        paid = int(record['paid_amount'] or 0)
        values = [offset - 1, f"{record['first_name']} {record['last_name']}", record['code'],
                  record['grade'] or '', record['class_name'] or '', record['year'], record['month'],
                  record['service_type'], amount, paid, amount - paid, record['status'],
                  record['payment_date'] or '', record['description'] or '']
        for column, value in enumerate(values, 1):
            ws_records.cell(offset, column, _excel_value(value))
    style_body(ws_records, 2, len(record_headers))
    for column, width in enumerate((7, 24, 14, 12, 16, 8, 10, 14, 14, 14, 14, 14, 13, 26), 1):
        ws_records.column_dimensions[get_column_letter(column)].width = width
    ws_records.freeze_panes = 'A2'
    ws_records.auto_filter.ref = f'A1:{get_column_letter(len(record_headers))}{max(1, ws_records.max_row)}'
    ws_records.page_setup.orientation = 'landscape'
    ws_records.page_setup.paperSize = ws_records.PAPERSIZE_A4
    ws_records.sheet_properties.pageSetUpPr.fitToPage = True
    ws_records.page_setup.fitToWidth = 1
    ws_records.page_setup.fitToHeight = 0
    ws_records.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    ws_records.print_title_rows = '1:1'

    debtors = [record for record in rows if int(record['amount'] or 0) > 0 and int(record['amount'] or 0) > int(record['paid_amount'] or 0)]
    ws_debtors = wb.create_sheet('بدهکاران')
    ws_debtors.sheet_view.rightToLeft = True
    debtor_headers = ['ردیف', 'دانش‌آموز', 'کد', 'کلاس', 'سال', 'ماه', 'مبلغ مصوب', 'پرداختی', 'مانده', 'وضعیت', 'تاریخ پرداخت']
    for column, header in enumerate(debtor_headers, 1):
        ws_debtors.cell(1, column, header)
    style_header(ws_debtors, 1, len(debtor_headers))
    for offset, record in enumerate(debtors, 2):
        amount = int(record['amount'] or 0)
        paid = int(record['paid_amount'] or 0)
        values = [offset - 1, f"{record['first_name']} {record['last_name']}", record['code'],
                  record['class_name'] or '', record['year'], record['month'], amount, paid, amount - paid,
                  record['status'], record['payment_date'] or '']
        for column, value in enumerate(values, 1):
            ws_debtors.cell(offset, column, _excel_value(value))
    style_body(ws_debtors, 2, len(debtor_headers))
    for column, width in enumerate((7, 24, 14, 16, 8, 10, 14, 14, 14, 14, 13), 1):
        ws_debtors.column_dimensions[get_column_letter(column)].width = width
    ws_debtors.freeze_panes = 'A2'
    ws_debtors.auto_filter.ref = f'A1:{get_column_letter(len(debtor_headers))}{max(1, ws_debtors.max_row)}'

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name='گزارش_سرویس_ماهانه.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _excel_value(value):
    if value is None:
        return ''
    if isinstance(value, str) and value[:1] in {'=', '+', '-', '@'}:
        return "'" + value
    return value


def monthly_service_print():
    filters = _service_filters()
    kind = request.args.get('kind', 'both')
    if kind not in ('requests', 'receipts', 'both'):
        kind = 'both'
    ids_raw = request.args.get('ids', '').strip()
    ids = [int(part) for part in ids_raw.split(',') if part.strip().isdigit()][:500] if ids_raw else []
    where, params = _service_where(filters)
    clauses = [where]
    if ids:
        clauses.append('ms.id IN (' + ','.join('?' for _ in ids) + ')')
        params = params + ids
    if kind == 'requests':
        clauses.append("COALESCE(ms.amount,0)>0 AND ms.status IN ('پرداخت نشده','جزئی')")
    elif kind == 'receipts':
        clauses.append('COALESCE(ms.paid_amount,0)>0')
    month_order_sql = 'CASE ms.month ' + ' '.join(
        f"WHEN '{month}' THEN {order}" for month, order in MONTH_ORDER.items()
    ) + ' ELSE 99 END'
    with get_db() as conn:
        rows = conn.execute(
            f'''SELECT ms.*, s.first_name, s.last_name, s.code,
                       COALESCE(NULLIF(TRIM(s.class_name),''),'ثبت نشده') AS class_name,
                       COALESCE(NULLIF(TRIM(COALESCE(t.first_name,'') || ' ' || COALESCE(t.last_name,'')),''),'بدون معلم') AS teacher_name,
                       COALESCE(NULLIF(TRIM(t.class_name),''), NULLIF(TRIM(s.class_name),''),'ثبت نشده') AS teacher_class
                FROM monthly_service ms
                JOIN students s ON ms.student_code=s.code
                LEFT JOIN teachers t ON TRIM(COALESCE(t.code,''))=TRIM(COALESCE(s.teacher_code,''))
                WHERE {' AND '.join(clauses)}
                ORDER BY COALESCE(NULLIF(TRIM(t.last_name),''),'بدون معلم') COLLATE NOCASE,
                         COALESCE(NULLIF(TRIM(t.first_name),''),'') COLLATE NOCASE,
                         COALESCE(NULLIF(TRIM(t.class_name),''), NULLIF(TRIM(s.class_name),''),'ثبت نشده') COLLATE NOCASE,
                         COALESCE(NULLIF(TRIM(s.last_name),''),'') COLLATE NOCASE,
                         COALESCE(NULLIF(TRIM(s.first_name),''),'') COLLATE NOCASE,
                         ms.year DESC, {month_order_sql}, ms.id''',
            params,
        ).fetchall()
        settings = _settings(conn)
        layouts = {
            'request': load_layout(conn, 'service_request', session.get('user_id')),
            'receipt': load_layout(conn, 'service_receipt', session.get('user_id')),
        }

    slips = []
    for row in rows:
        item = dict(row)
        amount = int(item['amount'] or 0)
        paid = int(item['paid_amount'] or 0)
        item['remaining'] = max(0, amount - paid)
        if kind in ('requests', 'both') and amount > 0 and item['remaining'] > 0 and item['status'] != 'کامل':
            slips.append({'kind': 'request', 'row': item})
        if kind in ('receipts', 'both') and paid > 0:
            slips.append({'kind': 'receipt', 'row': item})

    request_slips = [slip for slip in slips if slip['kind'] == 'request']
    receipt_slips = [slip for slip in slips if slip['kind'] == 'receipt']
    print_batches = []
    if request_slips:
        request_batch = build_slip_batch(request_slips, layouts['request'])
        request_batch['page_name'] = 'service-request-page'
        print_batches.append(request_batch)
    if receipt_slips:
        receipt_batch = build_slip_batch(receipt_slips, layouts['receipt'])
        receipt_batch['page_name'] = 'service-receipt-page'
        print_batches.append(receipt_batch)
    total_pages = sum(len(batch['pages']) for batch in print_batches)
    return render_template(
        'print_monthly_service.html',
        print_batches=print_batches,
        settings=settings,
        kind=kind,
        slip_count=len(slips),
        total_pages=total_pages,
        filters_description=_filters_description(filters),
    )


def register(app):
    app.add_url_rule('/update_student_service_fees', 'update_student_service_fees', update_student_service_fees, methods=['POST'])
    app.add_url_rule('/add_student_to_service', 'add_student_to_service', add_student_to_service, methods=['POST'])
    app.add_url_rule('/remove_student_service_fee/<student_code>', 'remove_student_service_fee', remove_student_service_fee, methods=['POST'])
    app.add_url_rule('/generate_monthly_services', 'generate_monthly_services', generate_monthly_services, methods=['POST'])
    app.add_url_rule('/delete_monthly_service/<int:id>', 'delete_monthly_service', delete_monthly_service, methods=['POST'])
    app.add_url_rule('/delete_zero_monthly_services', 'delete_zero_monthly_services', delete_zero_monthly_services, methods=['POST'])
    app.add_url_rule('/financial_help', 'financial_help', financial_help, methods=['GET', 'POST'])
    app.add_url_rule('/financial_help/<int:id>/delete', 'delete_financial_help', delete_financial_help, methods=['POST'])
    app.add_url_rule('/monthly_service', 'monthly_service', monthly_service, methods=['GET', 'POST'])
    app.add_url_rule('/monthly_service/settings', 'monthly_service_settings', save_monthly_service_settings, methods=['POST'])
    app.add_url_rule('/monthly_service/rebuild-status', 'rebuild_monthly_status', rebuild_monthly_status, methods=['POST'])
    app.add_url_rule('/monthly_service/export', 'monthly_service_export', monthly_service_export)
    app.add_url_rule('/monthly_service/print', 'monthly_service_print', monthly_service_print)
    app.add_url_rule('/payment_history/<student_code>', 'payment_history', payment_history)
    app.add_url_rule('/edit_payment/<int:id>', 'edit_payment', edit_payment, methods=['GET', 'POST'])
    app.add_url_rule('/reports/debtors_by_month', 'debtors_by_month', debtors_by_month)
