from __future__ import annotations

import json
import os
import re
import secrets
import sqlite3
import uuid
from datetime import datetime, timedelta
from io import BytesIO

from flask import abort, current_app, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

from ..database import auto_backup, get_db
from ..dates import (get_today_jalali, gregorian_to_jalali, normalize_jalali_date,
                     parse_jalali_date, today_string, tomorrow_string)
from ..security import audit, teacher_scope
from ..statistics_routes import statistics_dashboard


def _excel_cell(value):
    """Return a value safe for spreadsheet display, preventing formula injection."""
    if value is None:
        return ''
    if isinstance(value, str) and value[:1] in {'=', '+', '-', '@'}:
        return "'" + value
    return value


def _export_student_query():
    """Build the student export query from active dashboard-style filters."""
    clauses = []
    params: list[object] = []
    scope = teacher_scope()
    status = (request.args.get('status') or 'active').strip()
    status_clauses = {
        'active': "(s.status='فعال' OR s.status IS NULL OR TRIM(COALESCE(s.status,''))='')",
        'graduate': "s.status IN ('فارغ‌التحصیل','فارغ التحصیل')",
        'dropout': "s.status='ترک تحصیل'",
        'other': "s.status NOT IN ('فعال','فارغ‌التحصیل','فارغ التحصیل','ترک تحصیل')",
        'all': '1=1',
    }
    clauses.append(status_clauses.get(status, status_clauses['active']))
    if scope:
        clauses.append('s.teacher_code=?')
        params.append(scope)
    q = (request.args.get('q') or '').strip()[:100]
    if q:
        like = f'%{q}%'
        fields = ('first_name', 'last_name', 'code', 'grade', 'class_name', 'sida_class', 'gender')
        clauses.append('(' + ' OR '.join(f"COALESCE(s.{field},'') LIKE ?" for field in fields) + ')')
        params.extend([like] * len(fields))
    for key in ('grade', 'class_name', 'sida_class', 'gender'):
        value = (request.args.get(key) or '').strip()[:100]
        if value == '__empty__':
            clauses.append(f"NULLIF(TRIM(COALESCE(s.{key},'')),'') IS NULL")
        elif value:
            clauses.append(f's.{key}=?')
            params.append(value)
    teacher_code = (request.args.get('teacher_code') or '').strip()[:50]
    if not scope and teacher_code:
        if teacher_code == '__empty__':
            clauses.append("NULLIF(TRIM(COALESCE(s.teacher_code,'')),'') IS NULL")
        else:
            clauses.append('s.teacher_code=?')
            params.append(teacher_code)
    query = '''SELECT s.*, t.first_name AS teacher_first_name, t.last_name AS teacher_last_name
               FROM students s LEFT JOIN teachers t ON s.teacher_code=t.code
               WHERE ''' + ' AND '.join(clauses) + ' ORDER BY s.id DESC'
    return query, params


def _record_date(value: object) -> str:
    return normalize_jalali_date(str(value or '').strip(), required=True)


def _record_text(value: object, limit: int = 1000) -> str:
    return str(value or '').strip()[:limit]


def discipline():
    if request.method == 'POST':
        try:
            target_type = _record_text(request.form.get('target_type'), 30)
            target_code = _record_text(request.form.get('target_code'), 80)
            record_date = _record_date(request.form.get('date'))
            record_type = _record_text(request.form.get('type'), 30)
            points = int(request.form.get('points', '0'))
            description = _record_text(request.form.get('description'))
            if target_type not in {'دانش‌آموز', 'معلم'} or not target_code:
                raise ValueError('هدف انضباطی معتبر نیست.')
            if record_type not in {'مثبت', 'منفی'} or not 1 <= points <= 10:
                raise ValueError('نوع یا تعداد امتیاز معتبر نیست.')
            with get_db() as conn:
                table = 'students' if target_type == 'دانش‌آموز' else 'teachers'
                if not conn.execute(f'SELECT 1 FROM {table} WHERE code=?', (target_code,)).fetchone():
                    raise ValueError('فرد انتخاب‌شده پیدا نشد.')
            auto_backup()
            with get_db() as conn:
                conn.execute(
                    'INSERT INTO discipline(target_type,target_code,date,type,points,description) VALUES(?,?,?,?,?,?)',
                    (target_type, target_code, record_date, record_type, points, description),
                )
                conn.commit()
            audit(session.get('user_id'), 'create_discipline', 'discipline', target_code)
            flash('✅ امتیاز با موفقیت ثبت شد.', 'success')
        except (ValueError, TypeError):
            flash('❌ تاریخ، هدف یا امتیاز واردشده معتبر نیست.', 'danger')
        except Exception:
            current_app.logger.exception('management record operation failed')
            flash('❌ عملیات انجام نشد. اطلاعات واردشده را بررسی کنید.', 'danger')
        return redirect(url_for('discipline'))
    with get_db() as conn:
        records = conn.execute('SELECT * FROM discipline ORDER BY date DESC,id DESC').fetchall()
        students = conn.execute("SELECT first_name,last_name,code FROM students WHERE status='فعال' OR status IS NULL").fetchall()
        teachers = conn.execute('SELECT first_name,last_name,code FROM teachers').fetchall()
    return render_template('discipline.html', records=records, students=students, teachers=teachers, today=today_string())

def class_visit_report():
    if request.method == 'POST':
        try:
            teacher_code = _record_text(request.form.get('teacher_code'), 80)
            class_name = _record_text(request.form.get('class_name'), 80)
            record_date = _record_date(request.form.get('date'))
            report = _record_text(request.form.get('report'))
            if not teacher_code or not class_name or not report:
                raise ValueError('معلم، کلاس و گزارش الزامی هستند.')
            with get_db() as conn:
                if not conn.execute('SELECT 1 FROM teachers WHERE code=?', (teacher_code,)).fetchone():
                    raise ValueError('معلم انتخاب‌شده پیدا نشد.')
            auto_backup()
            with get_db() as conn:
                conn.execute(
                    'INSERT INTO class_visits(teacher_code,class_name,date,report) VALUES(?,?,?,?)',
                    (teacher_code, class_name, record_date, report),
                )
                conn.commit()
            audit(session.get('user_id'), 'create_class_visit', 'class_visit', teacher_code)
            flash('✅ گزارش بازدید ثبت شد.', 'success')
        except (ValueError, TypeError):
            flash('❌ تاریخ یا اطلاعات بازدید معتبر نیست.', 'danger')
        except Exception:
            current_app.logger.exception('management record operation failed')
            flash('❌ عملیات انجام نشد. اطلاعات واردشده را بررسی کنید.', 'danger')
        return redirect(url_for('class_visit_report'))
    with get_db() as conn:
        visits = conn.execute('SELECT * FROM class_visits ORDER BY date DESC,id DESC').fetchall()
        teachers = conn.execute('SELECT first_name,last_name,code FROM teachers').fetchall()
    return render_template('class_visit_report.html', visits=visits, teachers=teachers, today=today_string())

def teacher_performance():
    if request.method == 'POST':
        try:
            teacher_code = _record_text(request.form.get('teacher_code'), 80)
            student_code = _record_text(request.form.get('student_code'), 80)
            record_date = _record_date(request.form.get('date'))
            report = _record_text(request.form.get('report'))
            if not teacher_code or not student_code or not report:
                raise ValueError('معلم، دانش‌آموز و گزارش الزامی هستند.')
            with get_db() as conn:
                valid_pair = conn.execute(
                    """SELECT 1 FROM students s JOIN teachers t ON t.code=?
                       WHERE s.code=? AND (s.status='فعال' OR s.status IS NULL)""",
                    (teacher_code, student_code),
                ).fetchone()
                if not valid_pair:
                    raise ValueError('معلم یا دانش‌آموز انتخاب‌شده معتبر نیست.')
            auto_backup()
            with get_db() as conn:
                conn.execute(
                    'INSERT INTO teacher_performance(teacher_code,student_code,date,report) VALUES(?,?,?,?)',
                    (teacher_code, student_code, record_date, report),
                )
                conn.commit()
            audit(session.get('user_id'), 'create_teacher_performance', 'teacher_performance', student_code)
            flash('✅ گزارش عملکرد ثبت شد.', 'success')
        except (ValueError, TypeError):
            flash('❌ تاریخ یا اطلاعات عملکرد معتبر نیست.', 'danger')
        except Exception:
            current_app.logger.exception('management record operation failed')
            flash('❌ عملیات انجام نشد. اطلاعات واردشده را بررسی کنید.', 'danger')
        return redirect(url_for('teacher_performance'))
    with get_db() as conn:
        performances = conn.execute('''
            SELECT tp.*, s.first_name AS student_first, s.last_name AS student_last,
                   t.first_name AS teacher_first, t.last_name AS teacher_last
            FROM teacher_performance tp
            JOIN students s ON tp.student_code=s.code
            JOIN teachers t ON tp.teacher_code=t.code
            ORDER BY tp.date DESC,tp.id DESC
        ''').fetchall()
        students = conn.execute("SELECT first_name,last_name,code FROM students WHERE status='فعال' OR status IS NULL").fetchall()
        teachers = conn.execute('SELECT first_name,last_name,code FROM teachers').fetchall()
    return render_template('teacher_performance.html', performances=performances, students=students, teachers=teachers, today=today_string())

def edit_teacher_performance(id):
    with get_db() as conn:
        performance = conn.execute('SELECT * FROM teacher_performance WHERE id=?', (id,)).fetchone()
        if not performance:
            flash('⚠️ گزارش یافت نشد.', 'danger')
            return redirect(url_for('teacher_performance'))
        if request.method == 'POST':
            try:
                teacher_code = _record_text(request.form.get('teacher_code'), 80)
                student_code = _record_text(request.form.get('student_code'), 80)
                record_date = _record_date(request.form.get('date'))
                report = _record_text(request.form.get('report'))
                if not teacher_code or not student_code or not report:
                    raise ValueError('اطلاعات گزارش کامل نیست.')
                valid_pair = conn.execute(
                    """SELECT 1 FROM students s JOIN teachers t ON t.code=?
                       WHERE s.code=? AND (s.status='فعال' OR s.status IS NULL)""",
                    (teacher_code, student_code),
                ).fetchone()
                if not valid_pair:
                    raise ValueError('معلم یا دانش‌آموز انتخاب‌شده معتبر نیست.')
                auto_backup()
                conn.execute(
                    'UPDATE teacher_performance SET teacher_code=?,student_code=?,date=?,report=? WHERE id=?',
                    (teacher_code, student_code, record_date, report, id),
                )
                conn.commit()
                audit(session.get('user_id'), 'edit_teacher_performance', 'teacher_performance', id)
                flash('✅ گزارش عملکرد با موفقیت ویرایش شد.', 'success')
            except (ValueError, TypeError):
                flash('❌ تاریخ یا اطلاعات عملکرد معتبر نیست.', 'danger')
            except Exception:
                current_app.logger.exception('management record operation failed')
                flash('❌ عملیات انجام نشد. اطلاعات واردشده را بررسی کنید.', 'danger')
            return redirect(url_for('teacher_performance'))
        students = conn.execute("SELECT first_name,last_name,code FROM students WHERE status='فعال' OR status IS NULL").fetchall()
        teachers = conn.execute('SELECT first_name,last_name,code FROM teachers').fetchall()
    return render_template('edit_teacher_performance.html', performance=performance, students=students, teachers=teachers)

def delete_teacher_performance(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM teacher_performance WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_teacher_performance', 'teacher_performance', id)
        flash('✅ گزارش عملکرد حذف شد.', 'success')
    else:
        flash('گزارش عملکرد پیدا نشد.', 'warning')
    return redirect(url_for('teacher_performance'))

def delete_discipline(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM discipline WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_discipline', 'discipline', id)
        flash('رکورد انضباطی حذف شد.', 'success')
    else:
        flash('رکورد انضباطی پیدا نشد.', 'warning')
    return redirect(url_for('discipline'))


def delete_class_visit(id):
    auto_backup()
    with get_db() as conn:
        cursor = conn.execute('DELETE FROM class_visits WHERE id=?', (id,))
        conn.commit()
    if cursor.rowcount:
        audit(session.get('user_id'), 'delete_class_visit', 'class_visit', id)
        flash('گزارش بازدید کلاسی حذف شد.', 'success')
    else:
        flash('گزارش بازدید پیدا نشد.', 'warning')
    return redirect(url_for('class_visit_report'))


def export_excel_filtered():
    """Export the currently filtered students as a compact Persian print list."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    query, params = _export_student_query()
    with get_db() as conn:
        students = conn.execute(query, params).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'فهرست دانش‌آموزان'
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'A5'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.15, footer=0.15)
    ws.print_title_rows = '1:4'

    last_col = 9
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, 'فهرست دانش‌آموزان فیلترشده')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, f"{current_app.config['SCHOOL_NAME']}  |  تاریخ تهیه: {get_today_jalali()}  |  تعداد: {len(students)} نفر")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    filters = []
    labels = {'q':'جست‌وجو','grade':'پایه','class_name':'کلاس مدرسه','sida_class':'کلاس سیدا','gender':'جنسیت','teacher_code':'معلم'}
    for key, label in labels.items():
        value = (request.args.get(key) or '').strip()
        if value: filters.append(f'{label}: {value}')
    ws.cell(3, 1, ' | '.join(filters) if filters else 'فهرست دانش‌آموزان فعال')
    headers = ['ردیف', 'نام', 'نام خانوادگی', 'کد دانش‌آموزی', 'جنسیت', 'پایه', 'کلاس مدرسه', 'کلاس سیدا', 'معلم']
    for col, value in enumerate(headers, 1): ws.cell(4, col, value)
    for row_no, student in enumerate(students, 1):
        teacher = f"{student['teacher_first_name'] or ''} {student['teacher_last_name'] or ''}".strip()
        values = [row_no, student['first_name'] or '', student['last_name'] or '', student['code'] or '',
                  student['gender'] or 'ثبت نشده', student['grade'] or 'ثبت نشده', student['class_name'] or 'ثبت نشده',
                  student['sida_class'] or 'ثبت نشده', teacher or 'ثبت نشده']
        for col, value in enumerate(values, 1): ws.cell(row_no + 4, col, _excel_cell(value))

    navy = '145DA0'; light = 'EAF3FB'; border_color = 'C8D8E8'
    thin = Side(style='thin', color=border_color)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=navy); cell.font = Font(name='Vazirmatn', size=15, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for r in (2, 3):
        for cell in ws[r]:
            cell.font = Font(name='Vazirmatn', size=9, color='44546A', italic=(r == 3))
            cell.alignment = Alignment(horizontal='center', vertical='center')
    for cell in ws[4]:
        cell.fill = PatternFill('solid', fgColor=navy); cell.font = Font(name='Vazirmatn', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.font = Font(name='Vazirmatn', size=9, color='1F2937')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(bottom=thin)
        if row[0].row % 2 == 1:
            for cell in row: cell.fill = PatternFill('solid', fgColor='F7FAFC')
    widths = [8, 17, 24, 16, 12, 18, 18, 18, 24]
    for i, width in enumerate(widths, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 28; ws.row_dimensions[4].height = 25
    ws.auto_filter.ref = f'A4:{openpyxl.utils.get_column_letter(last_col)}{max(4, ws.max_row)}'
    ws.print_area = f'A1:{openpyxl.utils.get_column_letter(last_col)}{max(4, ws.max_row)}'
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='فهرست_دانش_آموزان_فیلترشده.xlsx')


def export_excel():
    import openpyxl
    from io import BytesIO
    with get_db() as conn:
        students = conn.execute('\n            SELECT \n                s.*,\n                t.first_name AS teacher_first_name,\n                t.last_name AS teacher_last_name\n            FROM students s\n            LEFT JOIN teachers t ON s.teacher_code = t.code\n            ORDER BY s.id DESC\n        ').fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['ID', 'Name', 'Last Name', 'Code', 'Birth Date', 'Grade', 'Class', 'Sida Class', 'Teacher', 'Father Name', 'Father Code', 'Mother Name', 'Mother Code', 'Father Phone', 'Mother Phone', 'Phone', 'Address', 'Admission Year', 'Admission Type', 'Religion', 'Insurance Type', 'Sisters', 'Brothers', 'Postal Code', 'Economic Decile', 'Medications', 'Father Job', 'Father Education', 'Mother Job', 'Mother Education', 'Repeat Grade 1', 'Repeat Grade 2', 'Repeat Grade 3', 'Repeat Grade 4', 'Repeat Grade 5', 'Repeat Grade 6', 'License Years', 'Child Order', 'Diseases', 'Vaccinations Complete', 'Drug Allergy', 'Behavioral Traits', 'Parents Alive', 'Parents Separated', 'Child Lives With', 'Parents Related', 'Martyr Quota', 'Financial Status', 'Father Birth Date', 'Father ID Serial', 'Father ID Serial Letter', 'Father ID Issue Place', 'Mother Birth Date', 'Mother ID Serial', 'Mother ID Serial Letter', 'Mother ID Issue Place', 'Student ID Serial', 'Student ID Serial Letter', 'Student ID Issue Place', 'Is Emdad Member', 'Is Behzisti Member', 'Entry Year Grade 1', 'Entry Year Grade 2', 'Entry Year Grade 3', 'Entry Year Grade 4', 'Entry Year Grade 5', 'Entry Year Grade 6', 'Photo', 'Student Birth Place', 'Gender', 'Student ID Issue Date', 'Father Birth Place', 'Father ID Issue Date', 'Mother Birth Place', 'Mother ID Issue Date', 'Service Fee', 'Status'])
    for s in students:
        teacher_first = s['teacher_first_name'] or ''
        teacher_last = s['teacher_last_name'] or ''
        teacher_name = f'{teacher_first} {teacher_last}'.strip()
        ws.append([s['id'], s['first_name'], s['last_name'], s['code'], s['birth_date'], s['grade'], s['class_name'], s['sida_class'], teacher_name, s['father_name'] or '', s['father_code'] or '', s['mother_name'] or '', s['mother_code'] or '', s['father_phone'] or '', s['mother_phone'] or '', s['parent_phone'] or '', s['address'] or '', s['admission_year'] or '', s['admission_type'] or '', s['religion'] or '', s['insurance_type'] or '', s['sisters_count'] or '0', s['brothers_count'] or '0', s['postal_code'] or '', s['economic_decile'] or '', s['medications'] or '', s['father_job'] or '', s['father_education'] or '', s['mother_job'] or '', s['mother_education'] or '', s['repeat_grade_1'] or '0', s['repeat_grade_2'] or '0', s['repeat_grade_3'] or '0', s['repeat_grade_4'] or '0', s['repeat_grade_5'] or '0', s['repeat_grade_6'] or '0', s['license_years'] or '', s['child_order'] or '0', s['diseases'] or '', s['vaccinations_complete'] or '0', s['drug_allergy'] or '0', s['behavioral_traits'] or '', s['parents_alive'] or '', s['parents_separated'] or '0', s['child_lives_with'] or '', s['parents_related'] or '0', s['martyr_quota'] or '0', s['financial_status'] or '', s['father_birth_date'] or '', s['father_id_serial'] or '', s['father_id_serial_letter'] or '', s['father_id_issue_place'] or '', s['mother_birth_date'] or '', s['mother_id_serial'] or '', s['mother_id_serial_letter'] or '', s['mother_id_issue_place'] or '', s['student_id_serial'] or '', s['student_id_serial_letter'] or '', s['student_id_issue_place'] or '', s['is_emdad_member'] or '0', s['is_behzisti_member'] or '0', s['entry_year_grade_1'] or '', s['entry_year_grade_2'] or '', s['entry_year_grade_3'] or '', s['entry_year_grade_4'] or '', s['entry_year_grade_5'] or '', s['entry_year_grade_6'] or '', s['photo'] or '', s['student_birth_place'] or '', s['gender'] or '', s['student_id_issue_date'] or '', s['father_birth_place'] or '', s['father_id_issue_date'] or '', s['mother_birth_place'] or '', s['mother_id_issue_date'] or '', s['service_fee'] or 0, s['status'] or ''])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='students.xlsx')

def group_action():
    action = request.form.get('action', '').strip()
    ids_str = request.form.get('ids', '').strip()
    if not ids_str:
        flash('⚠️ هیچ دانش\u200cآموزی انتخاب نشده است.', 'danger')
        return redirect(url_for('index'))
    ids = list(dict.fromkeys(int(x) for x in ids_str.split(',') if x.strip().isdigit()))
    if not ids:
        flash('⚠️ شناسه\u200cهای نامعتبر.', 'danger')
        return redirect(url_for('index'))
    if action == 'delete':
        return redirect(url_for('group_delete', ids=','.join(map(str, ids))))
    elif action == 'graduate':
        auto_backup()
        with get_db() as conn:
            placeholders = ','.join(['?'] * len(ids))
            conn.execute(f'UPDATE students SET status="فارغ\u200cالتحصیل" WHERE id IN ({placeholders})', ids)
            conn.commit()
        flash(f'✅ {len(ids)} دانش\u200cآموز فارغ\u200cالتحصیل شدند.', 'success')
    elif action == 'dropout':
        auto_backup()
        with get_db() as conn:
            placeholders = ','.join(['?'] * len(ids))
            conn.execute(f'UPDATE students SET status="ترک تحصیل" WHERE id IN ({placeholders})', ids)
            conn.commit()
        flash(f'✅ {len(ids)} دانش\u200cآموز ترک\u200cتحصیل شدند.', 'success')
    elif action == 'export':
        import openpyxl
        with get_db() as conn:
            placeholders = ','.join(['?'] * len(ids))
            students = conn.execute(
                f"SELECT * FROM students WHERE id IN ({placeholders}) AND (status='فعال' OR status IS NULL)", ids
            ).fetchall()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'دانش‌آموزان منتخب'
        ws.append(['ID', 'نام', 'نام خانوادگی', 'کد', 'تاریخ تولد', 'پایه', 'کلاس', 'کلاس سیدا',
                   'دهک اقتصادی', 'نام پدر', 'کد ملی پدر', 'نام مادر', 'کد ملی مادر', 'تلفن ولی',
                   'آدرس', 'وضعیت'])
        for s in students:
            ws.append([_excel_cell(value) for value in (
                s['id'], s['first_name'], s['last_name'], s['code'], s['birth_date'], s['grade'],
                s['class_name'], s['sida_class'], s['economic_decile'], s['father_name'],
                s['father_code'], s['mother_name'], s['mother_code'], s['parent_phone'],
                s['address'], s['status'] or 'فعال')])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name='selected_students.xlsx')
    return redirect(url_for('index'))

def group_delete():
    if request.method == 'POST':
        ids_str = request.form.get('ids', '').strip()
        reason = request.form.get('reason', '').strip()
        if not ids_str or not reason:
            flash('⚠️ اطلاعات ناقص است.', 'danger')
            return redirect(url_for('index'))
        ids = [int(x) for x in ids_str.split(',') if x.strip().isdigit()]
        if not ids:
            flash('⚠️ شناسه\u200cهای نامعتبر.', 'danger')
            return redirect(url_for('index'))
        status_map = {'فارغ\u200cالتحصیل': 'فارغ\u200cالتحصیل', 'ترک تحصیل': 'ترک تحصیل', 'سایر': 'سایر'}
        new_status = status_map.get(reason, 'سایر')
        auto_backup()
        with get_db() as conn:
            placeholders = ','.join(['?'] * len(ids))
            conn.execute(f'UPDATE students SET status=? WHERE id IN ({placeholders})', [new_status] + ids)
            conn.commit()
        audit(session.get('user_id'), 'group_student_status', 'student', ','.join(map(str, ids)), new_status)
        flash(f'✅ {len(ids)} دانش\u200cآموز به وضعیت "{new_status}" تغییر یافتند.', 'success')
        return redirect(url_for('index'))
    ids_str = request.args.get('ids', '').strip()
    if not ids_str:
        flash('⚠️ هیچ دانش\u200cآموزی انتخاب نشده است.', 'danger')
        return redirect(url_for('index'))
    ids = list(dict.fromkeys(int(x) for x in ids_str.split(',') if x.strip().isdigit()))
    if not ids:
        flash('⚠️ شناسه\u200cهای نامعتبر.', 'danger')
        return redirect(url_for('index'))
    with get_db() as conn:
        placeholders = ','.join(['?'] * len(ids))
        students = conn.execute(f'SELECT * FROM students WHERE id IN ({placeholders})', ids).fetchall()
    return render_template('group_delete.html', students=students, ids=ids)

def student_field_settings():
    """Manage extra student fields without changing the Python forms."""
    allowed_types = {'text', 'textarea', 'number', 'date', 'select', 'checkbox'}
    if request.method == 'POST':
        action = _record_text(request.form.get('action'), 20)
        try:
            with get_db() as conn:
                if action == 'create':
                    label = _record_text(request.form.get('label'), 120)
                    field_type = _record_text(request.form.get('field_type'), 20)
                    section = _record_text(request.form.get('section_label'), 80) or 'اطلاعات تکمیلی'
                    if not label or field_type not in allowed_types:
                        raise ValueError('عنوان و نوع فیلد را درست وارد کنید.')
                    options = []
                    for raw in str(request.form.get('options') or '').replace(',', '\\n').splitlines():
                        value = raw.strip()[:120]
                        if value and value not in options:
                            options.append(value)
                    if field_type == 'select' and not options:
                        raise ValueError('برای فیلد انتخابی حداقل یک گزینه بنویسید.')
                    sort_order = int(request.form.get('sort_order') or 100)
                    sort_order = max(0, min(sort_order, 9999))
                    conn.execute(
                        '''INSERT INTO student_custom_fields
                           (field_key,label,field_type,options_json,section_label,required,show_to_teacher,active,sort_order,created_at)
                           VALUES(?,?,?,?,?,?,?,?,?,datetime('now'))''',
                        (f'custom_{uuid.uuid4().hex}', label, field_type,
                         json.dumps(options, ensure_ascii=False), section,
                         1 if request.form.get('required') else 0,
                         1 if request.form.get('show_to_teacher') else 0, 1, sort_order),
                    )
                    conn.commit()
                    audit(session.get('user_id'), 'create_student_custom_field', 'student_custom_field', label)
                    flash('فیلد جدید برای فرم‌های ثبت و ویرایش ساخته شد.', 'success')
                elif action == 'toggle':
                    field_id = int(request.form.get('field_id') or 0)
                    conn.execute('UPDATE student_custom_fields SET active=1-active WHERE id=?', (field_id,))
                    conn.commit()
                    audit(session.get('user_id'), 'toggle_student_custom_field', 'student_custom_field', field_id)
                    flash('وضعیت نمایش فیلد تغییر کرد.', 'success')
                elif action == 'update':
                    field_id = int(request.form.get('field_id') or 0)
                    label = _record_text(request.form.get('label'), 120)
                    field_type = _record_text(request.form.get('field_type'), 20)
                    section = _record_text(request.form.get('section_label'), 80) or 'اطلاعات تکمیلی'
                    if not label or field_type not in allowed_types:
                        raise ValueError('عنوان یا نوع فیلد معتبر نیست.')
                    options = []
                    for raw in str(request.form.get('options') or '').replace(',', '\\n').splitlines():
                        value = raw.strip()[:120]
                        if value and value not in options:
                            options.append(value)
                    if field_type == 'select' and not options:
                        raise ValueError('برای فیلد انتخابی حداقل یک گزینه لازم است.')
                    conn.execute(
                        '''UPDATE student_custom_fields SET label=?,field_type=?,options_json=?,
                           section_label=?,required=?,show_to_teacher=?,sort_order=? WHERE id=?''',
                        (label, field_type, json.dumps(options, ensure_ascii=False), section,
                         1 if request.form.get('required') else 0,
                         1 if request.form.get('show_to_teacher') else 0,
                         max(0, min(int(request.form.get('sort_order') or 100), 9999)), field_id),
                    )
                    conn.commit()
                    audit(session.get('user_id'), 'update_student_custom_field', 'student_custom_field', field_id)
                    flash('تنظیمات فیلد ذخیره شد؛ اطلاعات دانش‌آموزان باقی ماند.', 'success')
                else:
                    raise ValueError('عملیات نامعتبر است.')
        except (ValueError, TypeError):
            flash('اطلاعات فیلد معتبر نیست؛ گزینه‌ها را بررسی کنید.', 'danger')
        except Exception:
            current_app.logger.exception('student custom field operation failed')
            flash('عملیات فیلد انجام نشد.', 'danger')
        return redirect(url_for('student_field_settings'))
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM student_custom_fields ORDER BY sort_order,id').fetchall()
    fields = []
    for row in rows:
        item = dict(row)
        try:
            options = json.loads(item.get('options_json') or '[]')
        except (TypeError, ValueError, json.JSONDecodeError):
            options = []
        item['options'] = options if isinstance(options, list) else []
        fields.append(item)
    return render_template('student_field_settings.html', fields=fields)


def excessive_absences():
    with get_db() as conn:
        absences = conn.execute('SELECT s.first_name, s.last_name, s.code, COUNT(a.id) AS abs_count FROM students s JOIN attendance_students a ON s.code=a.student_code WHERE a.status="غایب" AND (s.status="فعال" OR s.status IS NULL) GROUP BY s.code HAVING COUNT(a.id)>5 ORDER BY abs_count DESC').fetchall()
    return render_template('excessive_absences.html', absences=absences)



def register(app):
    app.add_url_rule('/discipline', endpoint='discipline', view_func=discipline, methods=['GET', 'POST'])
    app.add_url_rule('/discipline/<int:id>/delete', endpoint='delete_discipline', view_func=delete_discipline, methods=['POST'])
    app.add_url_rule('/class_visit_report', endpoint='class_visit_report', view_func=class_visit_report, methods=['GET', 'POST'])
    app.add_url_rule('/class_visit_report/<int:id>/delete', endpoint='delete_class_visit', view_func=delete_class_visit, methods=['POST'])
    app.add_url_rule('/teacher_performance', endpoint='teacher_performance', view_func=teacher_performance, methods=['GET', 'POST'])
    app.add_url_rule('/edit_teacher_performance/<int:id>', endpoint='edit_teacher_performance', view_func=edit_teacher_performance, methods=['GET', 'POST'])
    app.add_url_rule('/delete_teacher_performance/<int:id>', endpoint='delete_teacher_performance', view_func=delete_teacher_performance, methods=['POST'])
    app.add_url_rule('/export_excel', endpoint='export_excel', view_func=export_excel_filtered)
    app.add_url_rule('/group_action', endpoint='group_action', view_func=group_action, methods=['POST'])
    app.add_url_rule('/group_delete', endpoint='group_delete', view_func=group_delete, methods=['GET', 'POST'])
    app.add_url_rule('/statistics', endpoint='statistics', view_func=statistics_dashboard)
    app.add_url_rule('/reports/excessive_absences', endpoint='excessive_absences', view_func=excessive_absences)
    app.add_url_rule('/settings/student-fields', endpoint='student_field_settings',
                     view_func=student_field_settings, methods=['GET', 'POST'])
