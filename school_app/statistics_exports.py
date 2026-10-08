"""Read-only statistics exports from the same server-side filtered snapshot."""
from __future__ import annotations
import base64
import io
from pathlib import Path
from threading import BoundedSemaphore
from flask import current_app, render_template, send_file

_PDF_SLOTS=BoundedSemaphore(2)
STUDENT_COLUMNS=[('name','نام'),('code','کد'),('grade','پایه'),('class_name','کلاس'),('teacher_name','معلم'),('gender','جنسیت')]


def _metadata(ctx, detail):
    return [str(current_app.config.get('SCHOOL_NAME',''))+' — '+('فهرست دانش‌آموزان' if detail else 'گزارش فراوانی'),
        ctx['scope_label'],ctx['selection_label'],
        ('جزئیات: '+ctx['detail_label']+' · تعداد: '+str(ctx['detail_count']) if detail else 'حالت: '+ctx['view_options'][ctx['view']]+' · مبنای سهم: '+str(ctx['population'])+' نفر · انتخاب فعلی: '+str(ctx['selected_population'])+' نفر')+' · استخراج: '+ctx['extracted_at']]


def _cell(row,col):
    if row.get('row_kind')!='detail' and col in ('teacher','grade'):
        return row.get('label','') if row.get('label_column')==col else ''
    if col=='teacher':
        code=row.get('code','')
        return row.get('teacher_name','')+(f' ({code})' if code and code!='__unassigned__' else '')
    return row.get(col,'')


def excel(ctx,detail=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='دانش‌آموزان' if detail else 'آمار';ws.sheet_view.rightToLeft=True
    columns=STUDENT_COLUMNS if detail else [(c['key'],c['label']) for c in ctx['frequency_columns_view']]
    for line in _metadata(ctx,detail):
        ws.append([line]);ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(columns))
    ws.append([label for key,label in columns]);header=ws.max_row
    kinds={}
    if detail:
        for row in ctx['student_rows']:ws.append([row.get(key,'') for key,label in columns])
    else:
        for row in ctx['frequency_display_rows']:
            if row['row_kind']=='section':
                ws.append([row['label']]);ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(columns))
            else:ws.append([row['share']/100 if key=='share' else _cell(row,key) for key,label in columns])
            kinds[ws.max_row]=row['row_kind']
        if ctx['frequency_rows']:
            total=ctx['frequency_totals'];ws.append(['جمع کل' if key==ctx['totals_label_key'] else '' if key in ('teacher','grade') else total['share']/100 if key=='share' else total[key] for key,label in columns]);kinds[ws.max_row]='grand_total'
    last_data=ws.max_row
    ws.append(['دادهٔ زمان استخراج؛ جمع کل بدون شمارش دوبارهٔ جمع‌های جزئی است. اطلاعات را بدون جمع‌زدن دوبارهٔ سطرهای جمع تحلیل کنید.'])
    ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(columns))
    shades={'grade_total':'EAF2FB','teacher_total':'E8F5EF','section':'F0F4F7','grand_total':'D5E5ED'}
    border=Border(bottom=Side(style='hair',color='CBD5E1'))
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s' # names/codes can never become formulas
            cell.font=Font(name='Vazirmatn',size=11,bold=cell.row<=header or cell.row in kinds)
            cell.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
            if header<=cell.row<=last_data:cell.border=border
            if cell.row==header:cell.fill=PatternFill('solid',fgColor='174C63');cell.font=Font(name='Vazirmatn',size=11,bold=True,color='FFFFFF')
            elif cell.row in kinds:cell.fill=PatternFill('solid',fgColor=shades.get(kinds[cell.row],'FFFFFF'))
            if not detail and cell.row>header and cell.row<=last_data and columns[cell.column-1][0]=='share' and isinstance(cell.value,(int,float)):cell.number_format='0.0%'
    for index,(key,label) in enumerate(columns,1):ws.column_dimensions[get_column_letter(index)].width=30 if key in ('teacher','teacher_name','name','grade') else 20 if key=='code' else 14
    for r in range(1,header):ws.row_dimensions[r].height=32
    for r in range(header,ws.max_row+1):ws.row_dimensions[r].height=25
    ws.freeze_panes=f'A{header+1}';ws.print_title_rows=f'1:{header}';ws.print_area=f'A1:{get_column_letter(len(columns))}{ws.max_row}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation=ctx['orientation'];ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.oddFooter.center.text='صفحه &P از &N';ws.oddFooter.center.font='Arial,Regular';ws.oddFooter.center.size=9
    if detail:ws.auto_filter.ref=f'A{header}:{get_column_letter(len(columns))}{last_data}'
    output=io.BytesIO();wb.save(output);output.seek(0);return output


def pdf(ctx,detail=False):
    if not _PDF_SLOTS.acquire(timeout=1):raise ValueError('دو خروجی PDF در حال آماده‌سازی است؛ کمی بعد دوباره تلاش کنید.')
    try:
        try:from playwright.sync_api import sync_playwright
        except ImportError:raise ValueError('برای PDF وابستگی‌ها و Chromium را نصب کنید؛ یا از چاپ مرورگر ← ذخیره به PDF استفاده کنید.') from None
        root=Path(current_app.static_folder)
        font=base64.b64encode((root/'Vazirmatn-Regular.ttf').read_bytes()).decode()
        css=(root/'statistics.css').read_text()
        source=render_template('statistics_export.html',**ctx,detail=detail,embedded_font=font,embedded_css=css,metadata=_metadata(ctx,detail))
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True)
                try:
                    context=browser.new_context(java_script_enabled=False,offline=True)
                    context.route('**/*',lambda route:route.abort())
                    page=context.new_page();page.set_content(source,wait_until='load',timeout=25000)
                    page.evaluate('document.fonts.ready')
                    return io.BytesIO(page.pdf(print_background=True,prefer_css_page_size=True))
                finally:browser.close()
        except Exception as exc:
            current_app.logger.warning('Statistics PDF unavailable: %s',type(exc).__name__)
            raise ValueError('PDF آماده نشد. python -m playwright install chromium را اجرا کنید؛ جایگزین: چاپ مرورگر ← ذخیره به PDF.') from None
    finally:_PDF_SLOTS.release()


def export_statistics(ctx,fmt,detail=False):
    output=excel(ctx,detail) if fmt=='xlsx' else pdf(ctx,detail)
    return send_file(output,as_attachment=True,download_name=('statistics-students' if detail else 'statistics')+'.'+fmt,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if fmt=='xlsx' else 'application/pdf',max_age=0)
